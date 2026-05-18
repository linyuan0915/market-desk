from __future__ import annotations

import shutil
import subprocess
import importlib.util
import base64
import hashlib
import hmac
import os
import math
import csv
import io
import time
import threading
import uuid
from decimal import Decimal
from datetime import date, datetime, timedelta
import json
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from fastapi import Body, FastAPI, HTTPException, Request, Response
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles


ROOT_DIR = Path(__file__).resolve().parents[1]
STATIC_DIR = ROOT_DIR / "app" / "static"
OUTPUT_DIR = ROOT_DIR / "output"
WATCHLIST_PATH = OUTPUT_DIR / "watchlist.json"
SENTIMENT_HISTORY_PATH = OUTPUT_DIR / "sentiment_history.json"
SKILL_DIR = Path.home() / ".codex" / "skills" / "market-breadth-heatmap-skill"
GENERATOR = SKILL_DIR / "scripts" / "generate.py"
TEMPLATE = SKILL_DIR / "assets" / "heatmap_template.html"
DAILY_BRIEF_DIR = ROOT_DIR.parent / "每日简报"
DAILY_BRIEF_SKILL_DIR = DAILY_BRIEF_DIR / "skills" / "daily-market-brief"
DAILY_BRIEF_OUTPUT_DIR = DAILY_BRIEF_DIR / "output"
DAILY_BRIEF_BUILD_SCRIPT = DAILY_BRIEF_DIR / "tmp" / "docx" / "build_5_15_brief.py"
RSSCAST_URL = "https://app-cn.rsscast.io/api/mcp/v1/mcp"
MARKET_DATA_DB = os.environ.get("MARKET_DATA_DB", "market_data")
EASTMONEY_KAMT_URL = "https://push2.eastmoney.com/api/qt/kamt/get"
EASTMONEY_INDEX_FLOW_URL = "https://push2.eastmoney.com/api/qt/ulist.np/get"
EASTMONEY_LIMIT_UP_URL = "https://push2ex.eastmoney.com/getTopicZTPool"
EASTMONEY_LIMIT_DOWN_URL = "https://push2ex.eastmoney.com/getTopicDTPool"
FRED_CSV_URL = "https://fred.stlouisfed.org/graph/fredgraph.csv"
TZ = ZoneInfo("Asia/Hong_Kong")
INDUSTRY_CATEGORY_MAP = {}
INDEX_META = [
    {"code": "000001", "name": "上证指数", "db_code": "000001.SS"},
    {"code": "399001", "name": "深证成指", "db_code": "399001.SZ"},
    {"code": "399006", "name": "创业板指", "db_code": "399006.SZ"},
    {"code": "000300", "name": "沪深300", "db_code": "000300.SS"},
    {"code": "000905", "name": "中证500", "db_code": "000905.SS"},
    {"code": "000688", "name": "科创50", "db_code": "000688.SS"},
]
DEFAULT_WATCHLIST = [
    {"code": "300750", "name": "宁德时代", "category": "新能源"},
    {"code": "603986", "name": "兆易创新", "category": "科技成长"},
    {"code": "300394", "name": "天孚通信", "category": "科技成长"},
    {"code": "001393", "name": "N维通利", "category": "新股观察"},
    {"code": "603333", "name": "尚纬股份", "category": "交易观察"},
    {"code": "603897", "name": "长城科技", "category": "交易观察"},
]
DEFAULT_WATCHLIST_CATEGORY = "未分类"
DATA_REFRESH_BATCH_SIZE = 5
DATA_REFRESH_SLEEP_SECONDS = 1.2
CACHE_TTL_SECONDS = int(os.environ.get("API_CACHE_TTL_SECONDS", "900"))
SESSION_COOKIE_NAME = "market_desk_session"
YAHOO_MARKET_SYMBOLS = [
    {"code": "^HSI", "name": "恒生指数", "market": "港股指数", "yahoo": "^HSI"},
    {"code": "^HSCE", "name": "恒生国企指数", "market": "港股指数", "yahoo": "^HSCE"},
    {"code": "HSTECH.HK", "name": "恒生科技指数", "market": "港股指数", "yahoo": "HSTECH.HK"},
    {"code": "^GSPC", "name": "标普500", "market": "美股指数", "yahoo": "^GSPC"},
    {"code": "^IXIC", "name": "纳斯达克综合", "market": "美股指数", "yahoo": "^IXIC"},
    {"code": "^DJI", "name": "道琼斯", "market": "美股指数", "yahoo": "^DJI"},
    {"code": "^VIX", "name": "VIX恐慌指数", "market": "美股波动率", "yahoo": "^VIX"},
]
FRED_SERIES = {
    "DGS10": {"group": "利率", "name": "美国10Y国债", "unit": "%", "source": "FRED/DGS10"},
    "DGS2": {"group": "利率", "name": "美国2Y国债", "unit": "%", "source": "FRED/DGS2"},
    "DTWEXBGS": {"group": "汇率", "name": "美元指数", "unit": "", "source": "FRED/DTWEXBGS"},
    "DEXCHUS": {"group": "汇率", "name": "美元/人民币", "unit": "", "source": "FRED/DEXCHUS"},
}
FUTURES_MAIN_CONTRACTS = [
    {"symbol": "AU0", "name": "黄金", "group": "商品", "unit": "元/克"},
    {"symbol": "SC0", "name": "原油", "group": "商品", "unit": "元/桶"},
    {"symbol": "CU0", "name": "铜", "group": "商品", "unit": "元/吨"},
    {"symbol": "I0", "name": "铁矿石", "group": "商品", "unit": "元/吨"},
    {"symbol": "RB0", "name": "螺纹钢", "group": "商品", "unit": "元/吨"},
]

if GENERATOR.exists():
    spec = importlib.util.spec_from_file_location("market_breadth_generate", GENERATOR)
    if spec and spec.loader:
        skill_generate = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(skill_generate)
        INDUSTRY_CATEGORY_MAP = getattr(skill_generate, "INDUSTRY_CATEGORY_MAP", {})

RESPONSE_CACHE: dict[str, dict] = {}
RESPONSE_CACHE_LOCK = threading.Lock()
DATA_REFRESH_JOBS: dict[str, dict] = {}
DATA_REFRESH_LOCK = threading.Lock()

app = FastAPI(title="市场宽度鱼盆热力图")
STATIC_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/output", StaticFiles(directory=OUTPUT_DIR), name="output")


@app.on_event("startup")
def startup() -> None:
    STATIC_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if not (OUTPUT_DIR / "market_breadth_heatmap.png").exists() and (ROOT_DIR / "market_breadth_heatmap.png").exists():
        shutil.copy2(ROOT_DIR / "market_breadth_heatmap.png", OUTPUT_DIR / "market_breadth_heatmap.png")
    _ensure_market_schema()


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/auth/status")
def auth_status(request: Request) -> dict:
    return {"enabled": _auth_enabled(), "authenticated": _is_authenticated(request)}


@app.get("/api/health")
def health() -> dict:
    database = {"ok": False, "message": "unchecked"}
    try:
        _ensure_market_schema()
        connection = _market_db_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    SELECT COUNT(*) AS table_count
                    FROM information_schema.tables
                    WHERE table_schema = %s AND table_name = 'daily_data'
                    """,
                    (MARKET_DATA_DB,),
                )
                row = cursor.fetchone() or {}
            database = {"ok": True, "message": "connected", "daily_data_table": bool(row.get("table_count"))}
        finally:
            connection.close()
    except Exception as error:
        database = {"ok": False, "message": str(error)}
    return {
        "ok": database["ok"],
        "app": "market-desk",
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "database": database,
        "auth_enabled": _auth_enabled(),
    }


@app.post("/api/login")
def login(response: Response, payload: dict = Body(...)) -> dict:
    if not _auth_enabled():
        return {"ok": True, "enabled": False}
    password = str(payload.get("password") or "")
    if not hmac.compare_digest(password, os.environ.get("APP_PASSWORD", "")):
        raise HTTPException(status_code=401, detail="访问密码错误。")
    response.set_cookie(
        SESSION_COOKIE_NAME,
        _make_session_token(),
        httponly=True,
        secure=os.environ.get("APP_COOKIE_SECURE", "0") == "1",
        samesite="lax",
        max_age=60 * 60 * 24 * 14,
    )
    return {"ok": True, "enabled": True}


@app.post("/api/logout")
def logout(response: Response) -> dict:
    response.delete_cookie(SESSION_COOKIE_NAME)
    return {"ok": True}


@app.get("/api/status")
def status() -> dict:
    image = OUTPUT_DIR / "market_breadth_heatmap.png"
    raw = OUTPUT_DIR / "raw_data.json"
    aggregated = OUTPUT_DIR / "aggregated_data.json"
    stats = _breadth_stats(aggregated)
    return {
        "image_url": f"/output/{image.name}" if image.exists() else None,
        "image_path": str(image) if image.exists() else None,
        "raw_data_path": str(raw) if raw.exists() else None,
        "aggregated_data_path": str(aggregated) if aggregated.exists() else None,
        "generated_at": datetime.fromtimestamp(image.stat().st_mtime, TZ).isoformat(timespec="seconds") if image.exists() else None,
        "skill": str(SKILL_DIR),
        "stats": stats,
        "commentary": _market_commentary(stats),
    }


@app.get("/api/market-breadth/data")
def market_breadth_data() -> JSONResponse:
    aggregated = OUTPUT_DIR / "aggregated_data.json"
    if not aggregated.exists():
        return JSONResponse(
            {
                "categories": [],
                "dates": [],
                "data": {},
                "stats": None,
                "commentary": {
                    "conclusion": "市场宽度数据暂不可用。",
                    "analysis": ["云端暂未包含本地 market-breadth skill 生成文件。当前模块需要先在本地生成 output/aggregated_data.json，或后续改为云端可调用的数据源。"],
                },
                "children": {},
            }
        )
    with aggregated.open(encoding="utf-8") as file:
        payload = json.load(file)
    stats = _breadth_stats(aggregated)
    payload["stats"] = stats
    payload["commentary"] = _market_commentary(stats)
    payload["children"] = _secondary_industry_data(OUTPUT_DIR / "raw_data.json")
    return JSONResponse(payload)


@app.post("/api/refresh")
def refresh(request: Request) -> dict:
    try:
        result = generate_heatmap()
        _cache_delete_prefix("market-breadth")
        _cache_delete_prefix("desk")
        return {**result, **status()}
    except Exception as error:
        return {**status(), "ok": False, "message": f"市场宽度刷新暂不可用：{error}"}


@app.get("/api/daily-brief/status")
def daily_brief_status() -> JSONResponse:
    docx_path = _latest_daily_brief_docx()
    if not docx_path:
        return JSONResponse(_fallback_daily_brief_payload(generated=False))
    return JSONResponse(_daily_brief_payload(docx_path))


@app.post("/api/daily-brief/generate")
def generate_daily_brief(request: Request) -> JSONResponse:
    if not DAILY_BRIEF_BUILD_SCRIPT.exists():
        return JSONResponse(_fallback_daily_brief_payload(generated=True))
    completed = subprocess.run(
        ["python3", str(DAILY_BRIEF_BUILD_SCRIPT)],
        cwd=DAILY_BRIEF_DIR,
        capture_output=True,
        text=True,
        timeout=120,
    )
    if completed.returncode != 0:
        return JSONResponse(
            {"ok": False, "message": completed.stderr or completed.stdout, "skill": str(DAILY_BRIEF_SKILL_DIR)},
            status_code=500,
        )
    docx_path = _latest_daily_brief_docx()
    if not docx_path:
        return JSONResponse({"ok": False, "message": "构建完成但未找到输出 Word。"}, status_code=500)
    return JSONResponse({"ok": True, "build_log": completed.stdout, **_daily_brief_payload(docx_path)})


@app.get("/api/daily-brief/download")
def download_daily_brief() -> FileResponse:
    docx_path = _latest_daily_brief_docx()
    if not docx_path:
        raise HTTPException(status_code=404, detail="未找到每日行情简报 Word 输出。")
    return FileResponse(
        docx_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=docx_path.name,
    )


@app.get("/api/a-share/indices")
def a_share_indices(request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload("a-share-indices", _a_share_indices_payload, refresh=refresh))


def _a_share_indices_payload() -> dict:
    codes = [item["code"] for item in INDEX_META]
    meta = {item["code"]: item for item in INDEX_META}
    end_date = datetime.now(TZ).date()
    start_date = end_date - timedelta(days=90)
    quotes = _rsscast_call("StockIndexPriceQuery", {"codes": codes})
    history = _rsscast_call(
        "StockIndexKLineQuery",
        {"codes": codes, "startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
    )
    rows = [_enrich_quote(row, meta.get(row["code"], {}).get("name", row["code"])) for row in quotes]
    histories = _group_history(history, meta)
    payload = {
        "source": "RssCast MCP",
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "trade_time": _latest_time(rows),
        "indices": sorted(rows, key=lambda row: [item["code"] for item in INDEX_META].index(row["code"])),
        "histories": histories,
        "summary": _index_summary(rows, histories),
    }
    return payload


@app.get("/api/market-sentiment")
def market_sentiment(request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload("market-sentiment", _market_sentiment_payload, refresh=refresh))


def _market_sentiment_payload() -> dict:
    codes = [item["code"] for item in INDEX_META]
    meta = {item["code"]: item for item in INDEX_META}
    end_date = datetime.now(TZ).date()
    start_date = end_date - timedelta(days=90)
    index_quotes = _rsscast_call("StockIndexPriceQuery", {"codes": codes})
    index_history = _rsscast_call(
        "StockIndexKLineQuery",
        {"codes": codes, "startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
    )
    indices = [_enrich_quote(row, meta.get(row["code"], {}).get("name", row["code"])) for row in index_quotes]
    histories = _group_history(index_history, meta)
    watch_payload = _watchlist_payload(_load_watchlist())
    breadth_stats = _breadth_stats(OUTPUT_DIR / "aggregated_data.json")
    payload = _sentiment_payload(indices, histories, breadth_stats, watch_payload)
    _record_sentiment(payload)
    payload["p1"] = _sentiment_p1_payload(payload, breadth_stats)
    return payload


@app.get("/api/watchlist")
def get_watchlist(request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload("watchlist", lambda: _watchlist_payload(_load_watchlist()), refresh=refresh))


@app.post("/api/watchlist")
def add_watchlist_item(request: Request, payload: dict = Body(...)) -> JSONResponse:
    _require_auth(request)
    code = str(payload.get("code", "")).strip()
    name = str(payload.get("name", "")).strip()
    category = str(payload.get("category", "")).strip() or DEFAULT_WATCHLIST_CATEGORY
    if not code.isdigit() or len(code) != 6:
        raise HTTPException(status_code=400, detail="股票代码必须是 6 位数字。")
    items = _load_watchlist()
    existing = next((item for item in items if item["code"] == code), None)
    if existing:
        if name:
            existing["name"] = name
        existing["category"] = category
    else:
        items.append({"code": code, "name": name or code, "category": category})
    _save_watchlist(items)
    _invalidate_data_views()
    return JSONResponse(_watchlist_payload(items))


@app.delete("/api/watchlist/{code}")
def remove_watchlist_item(request: Request, code: str) -> JSONResponse:
    _require_auth(request)
    items = [item for item in _load_watchlist() if item["code"] != code]
    _save_watchlist(items)
    _invalidate_data_views()
    return JSONResponse(_watchlist_payload(items))


@app.get("/api/a-share/stock/{code}")
def stock_detail(code: str) -> JSONResponse:
    if not code.isdigit() or len(code) != 6:
        raise HTTPException(status_code=400, detail="股票代码必须是 6 位数字。")
    end_date = datetime.now(TZ).date()
    start_date = end_date - timedelta(days=365)
    quote_rows = _rsscast_call("StockPriceQuery", {"codes": [code]})
    history_rows = _rsscast_call(
        "StockKLineQuery",
        {"codes": [code], "startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
    )
    if not quote_rows and not history_rows:
        raise HTTPException(status_code=404, detail="未查询到该股票数据。")
    name = _watchlist_name(code)
    quote = _enrich_quote(quote_rows[0], name) if quote_rows else None
    history = _history_metrics(sorted(history_rows, key=lambda row: row.get("timeString", "")))
    return JSONResponse(
        {
            "source": "RssCast MCP",
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "quote": quote,
            "history": history,
            "signals": _stock_signals(quote, history),
        }
    )


@app.get("/api/data-center")
def data_center(request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload("data-center", _safe_data_center_payload, refresh=refresh))


@app.post("/api/data-refresh")
def data_refresh(request: Request) -> JSONResponse:
    return JSONResponse(_start_data_refresh_job())


@app.get("/api/data-refresh/{job_id}")
def data_refresh_job(job_id: str) -> JSONResponse:
    job = _get_data_refresh_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="未找到该数据更新任务。")
    return JSONResponse(job)


@app.get("/api/desk")
def morning_desk(request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload("desk", _morning_desk_payload, refresh=refresh))


@app.get("/api/desk-card/{key}")
def desk_card(key: str, request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload(f"desk-card:{key}", lambda: _desk_card_payload(key), refresh=refresh))


@app.get("/api/cross-market-risk")
def cross_market_risk(request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload("cross-market-risk", _safe_cross_market_risk_payload, refresh=refresh))


@app.get("/api/sector-funds")
def sector_funds(request: Request, kind: str = "industry", refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload(f"sector-funds:{kind}", lambda: _sector_funds_payload(kind), refresh=refresh))


@app.get("/api/fund-mainline")
def fund_mainline(request: Request, kind: str = "industry", refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload(f"fund-mainline:{kind}", lambda: _fund_mainline_payload(kind), refresh=refresh))


@app.get("/api/macro-commodities")
def macro_commodities(request: Request, refresh: bool = False) -> JSONResponse:
    return JSONResponse(_cached_payload("macro-commodities", _macro_commodities_payload, refresh=refresh))


def generate_heatmap() -> dict:
    if not GENERATOR.exists():
        raise RuntimeError(f"未找到 skill 生成脚本: {GENERATOR}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    cmd = [
        "python3",
        str(GENERATOR),
        "--mode",
        "all",
        "--output-dir",
        str(OUTPUT_DIR),
        "--width",
        "1080",
        "--height",
        "1440",
    ]
    completed = subprocess.run(cmd, cwd=SKILL_DIR, capture_output=True, text=True, timeout=120)
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr or completed.stdout)
    return {"ok": True, "message": completed.stdout}


def _cached_payload(key: str, builder, refresh: bool = False, ttl: int = CACHE_TTL_SECONDS) -> dict:
    now = time.time()
    with RESPONSE_CACHE_LOCK:
        entry = RESPONSE_CACHE.get(key)
        if not refresh and entry and now - entry["created_at"] <= ttl:
            payload = entry["payload"].copy()
            payload["cache"] = {"hit": True, "key": key, "created_at": entry["created_at_iso"]}
            return payload
    payload = builder()
    if not isinstance(payload, dict):
        return payload
    created_at = datetime.now(TZ).isoformat(timespec="seconds")
    with RESPONSE_CACHE_LOCK:
        RESPONSE_CACHE[key] = {"created_at": now, "created_at_iso": created_at, "payload": payload}
    payload = payload.copy()
    payload["cache"] = {"hit": False, "key": key, "created_at": created_at}
    return payload


def _cache_delete_prefix(prefix: str) -> None:
    with RESPONSE_CACHE_LOCK:
        for key in list(RESPONSE_CACHE):
            if key.startswith(prefix):
                RESPONSE_CACHE.pop(key, None)


def _invalidate_data_views() -> None:
    for prefix in (
        "watchlist",
        "market-sentiment",
        "data-center",
        "cross-market-risk",
        "a-share-indices",
        "sector-funds",
        "fund-mainline",
        "macro-commodities",
        "desk",
        "desk-card",
    ):
        _cache_delete_prefix(prefix)


def _auth_enabled() -> bool:
    return bool(os.environ.get("APP_PASSWORD"))


def _session_secret() -> str:
    return os.environ.get("APP_SESSION_SECRET") or os.environ.get("APP_PASSWORD") or "local-dev"


def _make_session_token() -> str:
    issued_at = str(int(time.time()))
    digest = hmac.new(_session_secret().encode(), issued_at.encode(), hashlib.sha256).digest()
    return f"{issued_at}.{base64.urlsafe_b64encode(digest).decode().rstrip('=')}"


def _is_authenticated(request: Request) -> bool:
    if not _auth_enabled():
        return True
    token = request.cookies.get(SESSION_COOKIE_NAME, "")
    try:
        issued_at, signature = token.split(".", 1)
        if int(time.time()) - int(issued_at) > 60 * 60 * 24 * 14:
            return False
    except Exception:
        return False
    expected = _make_signature(issued_at)
    return hmac.compare_digest(signature, expected)


def _make_signature(value: str) -> str:
    digest = hmac.new(_session_secret().encode(), value.encode(), hashlib.sha256).digest()
    return base64.urlsafe_b64encode(digest).decode().rstrip("=")


def _require_auth(request: Request) -> None:
    if not _is_authenticated(request):
        raise HTTPException(status_code=401, detail="需要登录后才能执行该操作。")


def _require_auth_for_refresh(request: Request, refresh: bool) -> None:
    if refresh:
        _require_auth(request)


def _start_data_refresh_job() -> dict:
    job_id = uuid.uuid4().hex[:12]
    job = {
        "job_id": job_id,
        "status": "queued",
        "progress": 0,
        "message": "数据更新任务已进入队列。",
        "started_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "updated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "tasks": [],
        "warnings": [],
        "result": None,
    }
    with DATA_REFRESH_LOCK:
        DATA_REFRESH_JOBS[job_id] = job
    thread = threading.Thread(target=_run_data_refresh_job, args=(job_id,), daemon=True)
    thread.start()
    return job


def _get_data_refresh_job(job_id: str) -> dict | None:
    with DATA_REFRESH_LOCK:
        job = DATA_REFRESH_JOBS.get(job_id)
        return json.loads(json.dumps(job, default=str, ensure_ascii=False)) if job else None


def _update_data_refresh_job(job_id: str, **updates) -> None:
    with DATA_REFRESH_LOCK:
        job = DATA_REFRESH_JOBS.get(job_id)
        if not job:
            return
        job.update(updates)
        job["updated_at"] = datetime.now(TZ).isoformat(timespec="seconds")


def _run_data_refresh_job(job_id: str) -> None:
    started_at = datetime.now(TZ)
    task_plan = [
        ("A股指数", _refresh_a_share_indices_to_sql, 25),
        ("A股自选池个股", _refresh_a_share_stocks_to_sql, 65),
        ("港美指数/VIX", _refresh_yahoo_markets_to_sql, 90),
    ]
    tasks = []
    warnings = []
    try:
        _update_data_refresh_job(job_id, status="running", progress=5, message="正在连接 market_data 数据库。")
        connection = _market_db_connection()
    except Exception as error:
        _update_data_refresh_job(job_id, status="failed", progress=100, message=f"无法连接数据库：{error}", warnings=[str(error)])
        return

    try:
        for task_name, runner, progress in task_plan:
            _update_data_refresh_job(job_id, status="running", progress=max(5, progress - 15), message=f"正在更新{task_name}。", tasks=tasks)
            result = runner(connection)
            tasks.append(result)
            warnings.extend(result.get("warnings", []))
            _update_data_refresh_job(job_id, progress=progress, message=f"{task_name}更新完成。", tasks=tasks, warnings=warnings[:20])
        coverage = _sql_coverage_snapshot(connection)
    except Exception as error:
        warnings.append(str(error))
        coverage = {}
    finally:
        connection.close()

    ok = all(task.get("ok", False) for task in tasks)
    inserted = sum(int(task.get("rows_upserted") or 0) for task in tasks)
    result = {
        "ok": ok,
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "duration_seconds": round((datetime.now(TZ) - started_at).total_seconds(), 1),
        "message": f"数据更新完成，写入/更新 {inserted} 行。" if ok else "数据更新部分完成，存在数据源失败。",
        "rows_upserted": inserted,
        "tasks": tasks,
        "warnings": warnings,
        "coverage": coverage,
    }
    _invalidate_data_views()
    _update_data_refresh_job(
        job_id,
        status="completed" if ok else "completed_with_warnings",
        progress=100,
        message=result["message"],
        tasks=tasks,
        warnings=warnings[:20],
        result=result,
        finished_at=datetime.now(TZ).isoformat(timespec="seconds"),
    )


def _rsscast_call(tool: str, arguments: dict) -> list[dict]:
    token = os.environ.get("RSSCAST_MCP_TOKEN")
    if not token:
        raise HTTPException(status_code=500, detail="缺少 RSSCAST_MCP_TOKEN 环境变量，无法调用 RssCast MCP。")
    payload = {"jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": tool, "arguments": arguments}}
    response = requests.post(
        RSSCAST_URL,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
            "Accept": "application/json, text/event-stream",
        },
        json=payload,
        timeout=30,
    )
    response.raise_for_status()
    message = _parse_mcp_sse(response.text)
    if "error" in message:
        raise HTTPException(status_code=502, detail=message["error"])
    content = message.get("result", {}).get("content", [])
    if not content:
        return []
    text = content[0].get("text", "[]")
    return json.loads(text)


def _mysql_identifier(name: str) -> str:
    if not name or not all(char.isalnum() or char == "_" for char in name):
        raise RuntimeError(f"不安全的 MySQL 标识符：{name}")
    return f"`{name}`"


def _mysql_connection(database: str | None = MARKET_DATA_DB):
    import pymysql

    kwargs = dict(
        host=os.environ.get("MARKET_DB_HOST", "127.0.0.1"),
        port=int(os.environ.get("MARKET_DB_PORT", "3306")),
        user=os.environ.get("MARKET_DB_USER", "root"),
        password=os.environ.get("MARKET_DB_PASSWORD", ""),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
    )
    if database:
        kwargs["database"] = database
    return pymysql.connect(**kwargs)


def _market_db_connection():
    return _mysql_connection(MARKET_DATA_DB)


def _ensure_market_schema() -> None:
    try:
        database_name = _mysql_identifier(MARKET_DATA_DB)
        connection = _mysql_connection(None)
    except Exception:
        return
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                f"CREATE DATABASE IF NOT EXISTS {database_name} CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci"
            )
            cursor.execute(f"USE {database_name}")
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS daily_data (
                  date DATE NOT NULL,
                  code VARCHAR(32) NOT NULL,
                  name VARCHAR(128) NOT NULL,
                  market VARCHAR(32) NOT NULL,
                  close DECIMAL(20, 6) NULL,
                  change_amount DECIMAL(20, 6) NULL,
                  change_pct DECIMAL(20, 6) NULL,
                  volume DECIMAL(28, 4) NULL,
                  amount DECIMAL(28, 4) NULL,
                  created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
                  updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                  PRIMARY KEY (date, code, market),
                  KEY idx_daily_data_market_date (market, date),
                  KEY idx_daily_data_code_date (code, date),
                  KEY idx_daily_data_date (date)
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
                """
            )
        connection.commit()
    finally:
        connection.close()


def _data_refresh_payload() -> dict:
    started_at = datetime.now(TZ)
    tasks = []
    _ensure_market_schema()
    try:
        connection = _market_db_connection()
    except Exception as error:
        return {
            "ok": False,
            "generated_at": started_at.isoformat(timespec="seconds"),
            "message": f"无法连接本地数据库：{error}",
            "tasks": [],
        }

    try:
        tasks.append(_refresh_a_share_indices_to_sql(connection))
        tasks.append(_refresh_a_share_stocks_to_sql(connection))
        tasks.append(_refresh_yahoo_markets_to_sql(connection))
        coverage = _sql_coverage_snapshot(connection)
    finally:
        connection.close()

    ok = all(task.get("ok", False) for task in tasks)
    inserted = sum(int(task.get("rows_upserted") or 0) for task in tasks)
    warnings = [message for task in tasks for message in task.get("warnings", [])]
    return {
        "ok": ok,
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "duration_seconds": round((datetime.now(TZ) - started_at).total_seconds(), 1),
        "message": f"数据更新完成，写入/更新 {inserted} 行。" if ok else "数据更新部分完成，存在数据源失败。",
        "rows_upserted": inserted,
        "tasks": tasks,
        "warnings": warnings,
        "coverage": coverage,
    }


def _refresh_a_share_indices_to_sql(connection) -> dict:
    codes = [item["code"] for item in INDEX_META]
    end_date = datetime.now(TZ).date()
    start_date = _refresh_start_date(connection, "A股指数", end_date, fallback_days=10)
    try:
        rows = _rsscast_call(
            "StockIndexKLineQuery",
            {"codes": codes, "startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
        )
        meta = {item["code"]: item["name"] for item in INDEX_META}
        suffix_map = {item["code"]: item.get("db_code", item["code"]) for item in INDEX_META}
        count = _upsert_market_rows(connection, _normalize_rsscast_history_rows(rows, meta, "A股指数", suffix_map))
        return {"name": "A股指数", "ok": True, "source": "RssCast StockIndexKLineQuery", "start_date": start_date.isoformat(), "end_date": end_date.isoformat(), "rows_upserted": count, "warnings": []}
    except Exception as error:
        return {"name": "A股指数", "ok": False, "source": "RssCast StockIndexKLineQuery", "rows_upserted": 0, "warnings": [str(error)]}


def _refresh_a_share_stocks_to_sql(connection) -> dict:
    end_date = datetime.now(TZ).date()
    symbols = _watchlist_a_share_symbols(connection)
    total = 0
    warnings = []
    grouped: dict[date, list[dict]] = {}
    for item in symbols:
        grouped.setdefault(item["start_date"], []).append(item)
    group_items = sorted(grouped.items(), key=lambda item: item[0], reverse=True)
    request_count = 0
    for start_date, group_symbols in group_items:
        for index in range(0, len(group_symbols), DATA_REFRESH_BATCH_SIZE):
            batch_symbols = group_symbols[index : index + DATA_REFRESH_BATCH_SIZE]
            plain_codes = [item["plain_code"] for item in batch_symbols]
            suffix_map = {item["plain_code"]: item["code"] for item in batch_symbols}
            name_map = {item["plain_code"]: item["name"] or item["code"] for item in batch_symbols}
            try:
                rows = _rsscast_call(
                    "StockKLineQuery",
                    {"codes": plain_codes, "startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
                )
                returned_codes = {str(row.get("code") or "") for row in rows}
                missing_codes = [code for code in plain_codes if code not in returned_codes]
                if missing_codes:
                    warnings.append(f"{','.join(missing_codes)}: RssCast 未返回K线数据")
                total += _upsert_market_rows(connection, _normalize_rsscast_history_rows(rows, name_map, "A股个股", suffix_map))
            except Exception as error:
                warnings.append(f"{','.join(plain_codes)}: {error}")
            request_count += 1
            if request_count < sum(math.ceil(len(items) / DATA_REFRESH_BATCH_SIZE) for items in grouped.values()):
                time.sleep(DATA_REFRESH_SLEEP_SECONDS)
    return {
        "name": "A股个股",
        "ok": not warnings,
        "source": "RssCast StockKLineQuery",
        "symbols": len(symbols),
        "scope": "自选池",
        "start_date": min((item["start_date"] for item in symbols), default=end_date).isoformat(),
        "end_date": end_date.isoformat(),
        "rows_upserted": total,
        "warnings": warnings[:8],
    }


def _refresh_yahoo_markets_to_sql(connection) -> dict:
    end_date = datetime.now(TZ).date()
    start_date = _refresh_start_date(connection, "港美指数/VIX", end_date, fallback_days=10, markets=("港股指数", "美股指数", "美股波动率"))
    total = 0
    warnings = []
    for item in YAHOO_MARKET_SYMBOLS:
        try:
            rows = _yahoo_chart_rows(item, start_date, end_date)
            total += _upsert_market_rows(connection, rows)
        except Exception as error:
            warnings.append(f"{item['code']}: {error}")
    return {
        "name": "港美指数/VIX",
        "ok": not warnings,
        "source": "Yahoo Finance chart",
        "symbols": len(YAHOO_MARKET_SYMBOLS),
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "rows_upserted": total,
        "warnings": warnings[:8],
    }


def _refresh_start_date(connection, market: str, end_date: date, fallback_days: int = 10, markets: tuple[str, ...] | None = None) -> date:
    with connection.cursor() as cursor:
        if markets:
            placeholders = ",".join(["%s"] * len(markets))
            cursor.execute(f"SELECT MAX(date) AS max_date FROM daily_data WHERE market IN ({placeholders})", markets)
        else:
            cursor.execute("SELECT MAX(date) AS max_date FROM daily_data WHERE market = %s", (market,))
        row = cursor.fetchone() or {}
    latest = row.get("max_date")
    if isinstance(latest, datetime):
        latest = latest.date()
    if isinstance(latest, date):
        return max(latest - timedelta(days=3), end_date - timedelta(days=30))
    return end_date - timedelta(days=fallback_days)


def _watchlist_a_share_symbols(connection) -> list[dict]:
    end_date = datetime.now(TZ).date()
    output = []
    for item in _load_watchlist():
        plain_code = str(item.get("code") or "").strip()
        if not plain_code.isdigit() or len(plain_code) != 6:
            continue
        db_code = _a_share_db_code(plain_code)
        latest = _latest_symbol_date(connection, db_code)
        if latest:
            start_date = max(latest - timedelta(days=3), end_date - timedelta(days=30))
        else:
            start_date = end_date - timedelta(days=365)
        output.append(
            {
                "code": db_code,
                "plain_code": plain_code,
                "name": item.get("name") or plain_code,
                "start_date": start_date,
            }
        )
    return output


def _a_share_db_code(plain_code: str) -> str:
    return f"{plain_code}.SS" if plain_code.startswith("6") else f"{plain_code}.SZ"


def _latest_symbol_date(connection, code: str) -> date | None:
    with connection.cursor() as cursor:
        cursor.execute("SELECT MAX(date) AS max_date FROM daily_data WHERE market = 'A股个股' AND code = %s", (code,))
        row = cursor.fetchone() or {}
    latest = row.get("max_date")
    if isinstance(latest, datetime):
        return latest.date()
    return latest if isinstance(latest, date) else None


def _normalize_rsscast_history_rows(rows: list[dict], name_map: dict[str, str], market: str, suffix_map: dict[str, str] | None = None) -> list[dict]:
    suffix_map = suffix_map or {}
    sorted_rows = sorted(rows, key=lambda item: (str(item.get("code") or ""), str(item.get("timeString") or "")))
    previous_close_by_code: dict[str, float] = {}
    output = []
    for row in sorted_rows:
        code = str(row.get("code") or "")
        db_code = suffix_map.get(code, code)
        trade_date = str(row.get("timeString") or "")[:10]
        if not code or not trade_date:
            continue
        close = row.get("close")
        previous_close = previous_close_by_code.get(code)
        change_amount = row.get("change_amount", row.get("change"))
        change_pct = row.get("change_pct")
        if close is not None and previous_close:
            if change_amount is None:
                change_amount = float(close) - previous_close
            if change_pct is None:
                change_pct = (float(close) / previous_close - 1) * 100
        if close is not None:
            previous_close_by_code[code] = float(close)
        output.append(
            {
                "date": trade_date,
                "code": db_code,
                "name": name_map.get(code) or row.get("name") or db_code,
                "market": market,
                "close": close,
                "change": change_amount,
                "change_pct": change_pct,
                "volume": row.get("volume"),
                "amount": row.get("amount"),
            }
        )
    return output


def _yahoo_chart_rows(item: dict, start_date: date, end_date: date) -> list[dict]:
    period1 = int(datetime.combine(start_date, datetime.min.time(), tzinfo=TZ).timestamp())
    period2 = int(datetime.combine(end_date + timedelta(days=1), datetime.min.time(), tzinfo=TZ).timestamp())
    response = requests.get(
        f"https://query1.finance.yahoo.com/v8/finance/chart/{item['yahoo']}",
        params={"period1": period1, "period2": period2, "interval": "1d"},
        headers={"User-Agent": "Mozilla/5.0"},
        timeout=20,
    )
    response.raise_for_status()
    payload = response.json()
    result = (payload.get("chart") or {}).get("result") or []
    if not result:
        raise RuntimeError(((payload.get("chart") or {}).get("error") or {}).get("description") or "Yahoo 未返回行情")
    chart = result[0]
    timestamps = chart.get("timestamp") or []
    quote = (chart.get("indicators") or {}).get("quote", [{}])[0]
    closes = quote.get("close") or []
    opens = quote.get("open") or []
    highs = quote.get("high") or []
    lows = quote.get("low") or []
    volumes = quote.get("volume") or []
    rows = []
    previous_close = None
    for index, timestamp in enumerate(timestamps):
        close = closes[index] if index < len(closes) else None
        if close is None:
            continue
        current_date = datetime.fromtimestamp(timestamp, TZ).date().isoformat()
        change = close - previous_close if previous_close else None
        change_pct = (close / previous_close - 1) * 100 if previous_close else None
        previous_close = close
        rows.append(
            {
                "date": current_date,
                "code": item["code"],
                "name": item["name"],
                "market": item["market"],
                "close": close,
                "change": change,
                "change_pct": change_pct,
                "volume": volumes[index] if index < len(volumes) else None,
                "amount": None,
                "open": opens[index] if index < len(opens) else None,
                "high": highs[index] if index < len(highs) else None,
                "low": lows[index] if index < len(lows) else None,
            }
        )
    return rows


def _upsert_market_rows(connection, rows: list[dict]) -> int:
    change_column = _daily_data_change_column(connection)
    values = [
        (
            row.get("date"),
            row.get("code"),
            row.get("name"),
            row.get("market"),
            row.get("close"),
            row.get("change_amount", row.get("change")),
            row.get("change_pct"),
            row.get("volume"),
            row.get("amount"),
        )
        for row in rows
        if row.get("date") and row.get("code")
    ]
    if not values:
        return 0
    with connection.cursor() as cursor:
        cursor.executemany(
            f"""
            INSERT INTO daily_data (date, code, name, market, close, {change_column}, change_pct, volume, amount)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                name = VALUES(name),
                market = VALUES(market),
                close = VALUES(close),
                {change_column} = VALUES({change_column}),
                change_pct = VALUES(change_pct),
                volume = VALUES(volume),
                amount = VALUES(amount)
            """,
            values,
        )
    connection.commit()
    return len(values)


def _daily_data_change_column(connection) -> str:
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT COLUMN_NAME
            FROM INFORMATION_SCHEMA.COLUMNS
            WHERE TABLE_SCHEMA = DATABASE()
              AND TABLE_NAME = 'daily_data'
              AND COLUMN_NAME IN ('change_amount', 'change')
            ORDER BY FIELD(COLUMN_NAME, 'change_amount', 'change')
            LIMIT 1
            """
        )
        row = cursor.fetchone() or {}
    column = row.get("COLUMN_NAME") if isinstance(row, dict) else row[0] if row else None
    return "change_amount" if column == "change_amount" else "`change`"


def _sql_coverage_snapshot(connection) -> dict:
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT market, COUNT(*) AS rows_count, COUNT(DISTINCT code) AS code_count, MIN(date) AS min_date, MAX(date) AS max_date
            FROM daily_data
            GROUP BY market
            ORDER BY market
            """
        )
        rows = cursor.fetchall()
    return {"markets": [_json_safe_row(row) for row in rows]}


def _data_center_payload() -> dict:
    _ensure_market_schema()
    try:
        connection = _market_db_connection()
    except Exception as error:
        return {
            "available": False,
            "source": f"MySQL/{MARKET_DATA_DB}",
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "message": f"无法连接本地数据库：{error}",
            "markets": [],
            "coverage": {},
            "quality": [],
        }
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    market,
                    COUNT(*) AS rows_count,
                    COUNT(DISTINCT code) AS code_count,
                    MIN(date) AS min_date,
                    MAX(date) AS max_date,
                    SUM(close IS NULL) AS missing_close,
                    SUM(change_pct IS NULL) AS missing_change_pct,
                    SUM(amount IS NULL) AS missing_amount
                FROM daily_data
                GROUP BY market
                ORDER BY rows_count DESC
                """
            )
            markets = cursor.fetchall()
            cursor.execute(
                """
                SELECT
                    COUNT(*) AS rows_count,
                    COUNT(DISTINCT code) AS code_count,
                    COUNT(DISTINCT market) AS market_count,
                    MIN(date) AS min_date,
                    MAX(date) AS max_date
                FROM daily_data
                """
            )
            coverage = cursor.fetchone() or {}
            cursor.execute(
                """
                SELECT date, market, COUNT(*) AS rows_count, COUNT(DISTINCT code) AS code_count
                FROM daily_data
                WHERE date >= (SELECT DATE_SUB(MAX(date), INTERVAL 10 DAY) FROM daily_data)
                GROUP BY date, market
                ORDER BY date DESC, market
                """
            )
            recent = cursor.fetchall()
            cursor.execute(
                """
                SELECT code, name, market, COUNT(*) AS days_count, MIN(date) AS min_date, MAX(date) AS max_date
                FROM daily_data
                GROUP BY code, name, market
                HAVING days_count < 10
                ORDER BY days_count ASC, code
                LIMIT 20
                """
            )
            sparse = cursor.fetchall()
    finally:
        connection.close()

    coverage = _json_safe_row(coverage)
    coverage = {
        "rows_count": int(coverage.get("rows_count") or 0),
        "code_count": int(coverage.get("code_count") or 0),
        "market_count": int(coverage.get("market_count") or 0),
        "min_date": coverage.get("min_date"),
        "max_date": coverage.get("max_date"),
    }
    quality = []
    for row in markets:
        days = _calendar_days(row.get("min_date"), row.get("max_date"))
        avg_rows = (row["rows_count"] / days) if days else 0
        quality.append(
            {
                "market": row["market"],
                "rows_count": row["rows_count"],
                "code_count": row["code_count"],
                "date_range": f"{row['min_date']} 至 {row['max_date']}",
                "avg_rows_per_calendar_day": round(avg_rows, 1),
                "missing_total": int((row.get("missing_close") or 0) + (row.get("missing_change_pct") or 0) + (row.get("missing_amount") or 0)),
                "status": _coverage_status(row),
            }
        )
    a_stock = next((row for row in markets if row["market"] == "A股个股"), None)
    suggestions = _data_center_suggestions(markets, a_stock)
    return {
        "available": True,
        "source": f"MySQL/{MARKET_DATA_DB}.daily_data",
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "coverage": coverage,
        "markets": [_json_safe_row(row) for row in markets],
        "quality": quality,
        "recent_calendar": [_json_safe_row(row) for row in recent],
        "sparse_symbols": [_json_safe_row(row) for row in sparse],
        "suggestions": suggestions,
    }


def _safe_data_center_payload() -> dict:
    try:
        return _data_center_payload()
    except Exception as error:
        return {
            "available": False,
            "source": f"MySQL/{MARKET_DATA_DB}.daily_data",
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "message": f"数据中心诊断失败：{error}",
            "coverage": {"rows_count": 0, "code_count": 0, "market_count": 0, "min_date": None, "max_date": None},
            "markets": [],
            "quality": [],
            "recent_calendar": [],
            "sparse_symbols": [],
            "suggestions": ["请确认云端 daily_data 表结构已自动创建，并执行“一键更新数据”。"],
        }


def _morning_desk_payload() -> dict:
    generated_at = datetime.now(TZ).isoformat(timespec="seconds")
    cross = _cross_market_risk_payload()
    macro = _macro_commodities_payload()
    funds = _fund_mainline_payload("industry")
    data_center = _data_center_payload()
    try:
        sentiment_payload = _market_sentiment_payload()
    except Exception as error:
        sentiment_payload = {"available": False, "score": None, "tag": "情绪不可用", "conclusion": f"市场情绪暂不可用：{error}"}
    try:
        watch = _watchlist_payload(_load_watchlist())
    except Exception as error:
        watch = {"items": [], "summary": {"conclusion": f"自选池暂不可用：{error}"}}

    cards = [
        _desk_card_cross(cross),
        _desk_card_sentiment(sentiment_payload),
        _desk_card_funds(funds),
        _desk_card_macro(macro),
        _desk_card_watch(watch),
    ]
    score_values = [card["score"] for card in cards if isinstance(card.get("score"), (int, float))]
    desk_score = round(sum(score_values) / len(score_values)) if score_values else 50
    conclusion = _desk_conclusion(desk_score, cards)
    return {
        "available": True,
        "source": "SQL + RssCast + market-breadth + Eastmoney/THS/AkShare + FRED/Yahoo",
        "generated_at": generated_at,
        "score": desk_score,
        "conclusion": conclusion,
        "cards": cards,
        "data_freshness": _desk_data_freshness(data_center),
        "details": {
            "cross_market": cross,
            "sentiment": sentiment_payload,
            "funds": funds,
            "macro": macro,
            "watchlist": watch,
        },
    }


def _desk_card_payload(key: str) -> dict:
    if key == "cross":
        return _desk_card_cross(_safe_cross_market_risk_payload())
    if key == "sentiment":
        try:
            return _desk_card_sentiment(_market_sentiment_payload())
        except Exception as error:
            return _desk_card_sentiment({"score": None, "tag": "情绪不可用", "conclusion": f"市场情绪暂不可用：{error}"})
    if key == "funds":
        return _desk_card_funds(_fund_mainline_payload("industry"))
    if key == "macro":
        return _desk_card_macro(_macro_commodities_payload())
    if key == "watch":
        try:
            return _desk_card_watch(_watchlist_payload(_load_watchlist()))
        except Exception as error:
            return _desk_card_watch({"items": [], "summary": {"conclusion": f"自选池暂不可用：{error}"}})
    if key == "freshness":
        return {
            "key": "freshness",
            "title": "数据新鲜度",
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "data_freshness": _desk_data_freshness(_safe_data_center_payload()),
        }
    raise HTTPException(status_code=404, detail="未知驾驶舱分块。")


def _desk_card_cross(cross: dict) -> dict:
    risk = cross.get("risk") or {}
    risk_score = risk.get("score")
    score = 100 - risk_score if isinstance(risk_score, (int, float)) else None
    return {
        "key": "cross",
        "title": "外围风险",
        "state": risk.get("state") or "数据不足",
        "score": score,
        "conclusion": (cross.get("summary") or {}).get("conclusion") or "外围风险暂不可用。",
        "why": "隔夜港股、美股和 VIX 会影响 A 股开盘风险偏好；风险分越高，早盘越需要先看权重承接。",
        "action": "高压时优先观察沪深300、上证指数和北向/主力资金是否同步走弱。",
        "source": cross.get("source") or "-",
    }


def _desk_card_sentiment(sentiment: dict) -> dict:
    return {
        "key": "sentiment",
        "title": "市场情绪",
        "state": sentiment.get("tag") or sentiment.get("label") or "数据不足",
        "score": sentiment.get("score"),
        "conclusion": sentiment.get("conclusion") or "市场情绪暂不可用。",
        "why": "情绪温度计把指数、宽度、量能、风格和自选池合成一个交易状态，适合判断今天是加风险还是降风险。",
        "action": "分数高但行业宽度不扩散时，避免追高单一主线；分数低但宽度企稳时，关注修复窗口。",
        "source": sentiment.get("source") or "-",
    }


def _desk_card_funds(funds: dict) -> dict:
    leader = (funds.get("items") or [{}])[0]
    score = _clamp(50 + float(leader.get("net5_yi") or 0) * 3 + int(leader.get("streak") or 0) * 8) if leader else None
    return {
        "key": "funds",
        "title": "资金主线",
        "state": leader.get("name") or "暂无主线",
        "score": round(score) if score is not None else None,
        "conclusion": (funds.get("summary") or {}).get("conclusion") or "资金主线暂不可用。",
        "why": "阶段资金榜比单日涨跌更适合识别资金是否连续集中；连续阶段为正说明主线更有持续性。",
        "action": "优先跟踪近3日与近5日同时净流入的行业，再看是否有概念扩散和成交额配合。",
        "source": funds.get("source") or "-",
    }


def _desk_card_macro(macro: dict) -> dict:
    analysis = (macro.get("summary") or {}).get("analysis") or []
    pressure = len(analysis)
    score = _clamp(70 - pressure * 10)
    return {
        "key": "macro",
        "title": "宏观压力",
        "state": "压力偏高" if pressure >= 2 else "压力中性" if pressure == 1 else "压力较低",
        "score": round(score),
        "conclusion": (macro.get("summary") or {}).get("conclusion") or "宏观商品暂不可用。",
        "why": "美债、美元、人民币和商品决定估值压力、外资风险偏好和资源品交易线索。",
        "action": analysis[0] if analysis else "继续观察美元、美债、人民币与工业品是否形成同向压力。",
        "source": macro.get("source") or "-",
    }


def _desk_card_watch(watch: dict) -> dict:
    rows = watch.get("items") or []
    pct_values = [float(row.get("change_pct_display") or 0) for row in rows]
    avg = sum(pct_values) / len(pct_values) if pct_values else 0
    up_ratio = sum(1 for value in pct_values if value > 0) / len(pct_values) if pct_values else 0.5
    score = _clamp(50 + avg * 10 + (up_ratio - 0.5) * 30)
    return {
        "key": "watch",
        "title": "自选异动",
        "state": "自选偏强" if score >= 60 else "自选偏弱" if score < 45 else "自选分化",
        "score": round(score),
        "conclusion": (watch.get("summary") or {}).get("conclusion") or "自选池暂不可用。",
        "why": "自选池反映你的真实关注资产，能补足全市场指标和个人持仓/研究方向之间的差异。",
        "action": "自选强于大盘时看是否有可持续主线；自选弱于大盘时检查是否集中在退潮方向。",
        "source": watch.get("source") or "-",
    }


def _desk_conclusion(score: int, cards: list[dict]) -> dict:
    weakest = min((card for card in cards if card.get("score") is not None), key=lambda item: item["score"], default={})
    strongest = max((card for card in cards if card.get("score") is not None), key=lambda item: item["score"], default={})
    if score >= 65:
        title = "早盘环境偏积极，可以围绕资金主线寻找扩散。"
    elif score >= 50:
        title = "早盘环境中性偏分化，先看主线能否延续。"
    else:
        title = "早盘环境偏谨慎，先控风险再看修复。"
    analysis = [
        f"综合分 {score}/100，最强信号来自{strongest.get('title', '-')}，最弱约束来自{weakest.get('title', '-')}。",
        "页面刷新会优先更新本地 SQL 的 A 股、港美指数和 VIX，再读取在线资金、宏观和情绪接口。",
        "如果某个数据源失败，卡片会保留来源和原因，不用旧数据冒充新判断。",
    ]
    return {"title": title, "analysis": analysis}


def _desk_data_freshness(data_center: dict) -> list[dict]:
    rows = []
    for row in data_center.get("markets") or []:
        rows.append(
            {
                "market": row.get("market"),
                "max_date": row.get("max_date"),
                "code_count": row.get("code_count"),
                "rows_count": row.get("rows_count"),
            }
        )
    return rows


def _cross_market_risk_payload() -> dict:
    _ensure_market_schema()
    try:
        connection = _market_db_connection()
    except Exception as error:
        return {"available": False, "source": f"MySQL/{MARKET_DATA_DB}", "message": f"无法连接本地数据库：{error}"}
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT date, code, name, market, close, change_pct, amount
                FROM daily_data
                WHERE market IN ('A股指数', '港股指数', '美股指数', '美股波动率')
                  AND date >= (SELECT DATE_SUB(MAX(date), INTERVAL 45 DAY) FROM daily_data)
                ORDER BY date ASC
                """
            )
            rows = cursor.fetchall()
    finally:
        connection.close()
    if not rows:
        return {
            "available": False,
            "source": f"MySQL/{MARKET_DATA_DB}.daily_data",
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "trade_date": None,
            "risk": {"score": None, "state": "等待数据"},
            "items": [],
            "histories": {},
            "summary": {"conclusion": "跨市场风险等待本地数据更新。", "analysis": ["云端数据库已连接，但 daily_data 暂无 A 股、港股、美股或 VIX 数据。请点击“一键更新数据”补齐数据。"]},
        }
    latest_by_code = {}
    history_by_code: dict[str, list[dict]] = {}
    for row in rows:
        code = row["code"]
        safe = _json_safe_row(row)
        history_by_code.setdefault(code, []).append(safe)
        latest_by_code[code] = safe
    latest = sorted(latest_by_code.values(), key=lambda row: (row["market"], row["code"]))
    risk = _risk_score_from_cross_market(latest, history_by_code)
    return {
        "available": True,
        "source": f"MySQL/{MARKET_DATA_DB}.daily_data",
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "trade_date": max((row["date"] for row in latest), default=None),
        "risk": risk,
        "items": latest,
        "histories": history_by_code,
        "summary": _cross_market_summary(risk, latest),
    }


def _safe_cross_market_risk_payload() -> dict:
    try:
        return _cross_market_risk_payload()
    except Exception as error:
        return {
            "available": False,
            "source": f"MySQL/{MARKET_DATA_DB}.daily_data",
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "trade_date": None,
            "risk": {"score": None, "state": "风险面板暂不可用"},
            "items": [],
            "histories": {},
            "summary": {
                "conclusion": "跨市场风险暂不可用。",
                "analysis": [f"读取云端 SQL 时发生异常：{error}", "请确认 daily_data 表已自动创建，并点击“一键更新数据”。"],
            },
        }


def _sector_funds_payload(kind: str) -> dict:
    kind_map = {
        "industry": ("行业", "行业资金流"),
        "concept": ("概念", "概念资金流"),
        "region": ("地域", "地域资金流"),
        "style": ("风格", None),
    }
    label, sector_type = kind_map.get(kind, kind_map["industry"])
    if not sector_type:
        return {
            "available": False,
            "source": "东方财富板块资金流",
            "kind": kind,
            "label": label,
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "summary": {"conclusion": "风格资金流暂未接入。", "analysis": ["东方财富/akshare 当前稳定封装覆盖行业、概念、地域资金流；风格维度缺少同口径主力资金字段，暂不混用风格指数或自定义标签替代。"]},
            "items": [],
        }
    try:
        import akshare as ak

        df = ak.stock_sector_fund_flow_rank(indicator="今日", sector_type=sector_type)
        rows = _normalize_sector_fund_rows(df, label)
        source = "东方财富板块资金流 / akshare"
    except Exception as error:
        rows, source = _sector_funds_fallback(kind, label, str(error))
        if not rows:
            return {
                "available": False,
                "source": source,
                "kind": kind,
                "label": label,
                "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
                "summary": {"conclusion": f"{label}主力资金暂不可用。", "analysis": [f"数据接口当前未返回有效数据：{error}"]},
                "items": [],
            }
    return {
        "available": True,
        "source": source,
        "kind": kind,
        "label": label,
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "summary": _sector_funds_summary(rows, label),
        "items": rows,
    }


def _fund_mainline_payload(kind: str) -> dict:
    if kind not in ("industry", "concept"):
        return {
            "available": False,
            "source": "东方财富历史板块资金流 / 同花顺阶段资金榜 / akshare",
            "kind": kind,
            "label": "地域" if kind == "region" else "风格",
            "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "summary": {"conclusion": "该维度暂不支持连续主线追踪。", "analysis": ["当前只对行业、概念做 3 日/5 日连续主力净流入追踪。"]},
            "items": [],
        }
    label = "行业" if kind == "industry" else "概念"
    rows, source, errors = _fund_mainline_from_stage_ranks(kind)
    if not rows:
        rows, source, errors = _fund_mainline_from_histories(kind, errors)
    rows = sorted(rows, key=lambda row: (row.get("streak", 0), row.get("net5_yi", 0), row.get("net3_yi", 0)), reverse=True)
    available = bool(rows)
    return {
        "available": available,
        "source": source,
        "kind": kind,
        "label": label,
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "summary": _fund_mainline_summary(rows, label, errors),
        "items": rows,
        "errors": errors[:5],
    }


def _fund_mainline_from_stage_ranks(kind: str) -> tuple[list[dict], str, list[str]]:
    import akshare as ak

    label = "行业" if kind == "industry" else "概念"
    fetcher = ak.stock_fund_flow_industry if kind == "industry" else ak.stock_fund_flow_concept
    errors = []
    stage_frames = {}
    for symbol, key in (("3日排行", "net3_yi"), ("5日排行", "net5_yi"), ("10日排行", "net10_yi")):
        try:
            stage_frames[key] = fetcher(symbol=symbol)
        except Exception as error:
            errors.append(f"同花顺{label}{symbol}: {error}")
    if "net3_yi" not in stage_frames or "net5_yi" not in stage_frames:
        return [], f"同花顺{label}阶段资金榜 / akshare", errors

    merged: dict[str, dict] = {}
    for key, df in stage_frames.items():
        rows = _normalize_ths_stage_fund_rows(df, kind, key)
        for row in rows:
            name = row["name"]
            current = merged.setdefault(
                name,
                {
                    "name": name,
                    "kind": kind,
                    "net3_yi": None,
                    "net5_yi": None,
                    "net10_yi": None,
                    "streak": 0,
                    "latest_date": "阶段排行",
                    "latest_net_yi": None,
                    "change_pct": None,
                    "history": [],
                },
            )
            current[key] = row[key]
            current["change_pct"] = row.get("change_pct")
    output = []
    for row in merged.values():
        net3 = row.get("net3_yi")
        net5 = row.get("net5_yi")
        net10 = row.get("net10_yi")
        if net3 is None or net5 is None:
            continue
        row["net3_yi"] = round(float(net3), 2)
        row["net5_yi"] = round(float(net5), 2)
        row["latest_net_yi"] = row["net3_yi"]
        row["streak"] = int(row["net3_yi"] > 0) + int(row["net5_yi"] > 0) + int((net10 or 0) > 0)
        row["history"] = [
            {"date": "近3日", "main_net_yi": row["net3_yi"], "change_pct": row.get("change_pct")},
            {"date": "近5日", "main_net_yi": row["net5_yi"], "change_pct": row.get("change_pct")},
        ]
        row["continuity_label"] = f"{row['streak']} 个阶段为正"
        if row["net3_yi"] > 0 and row["net5_yi"] > 0:
            output.append(row)
    return output[:80], f"同花顺{label}3/5/10日阶段资金榜 / akshare", errors


def _normalize_ths_stage_fund_rows(df, kind: str, key: str) -> list[dict]:
    rows = []
    for _, item in df.iterrows():
        name = str(item.get("行业") or item.get("概念") or "").strip()
        if not name:
            continue
        rows.append(
            {
                "name": name,
                "kind": kind,
                key: round(_to_float(item.get("净额")), 2),
                "change_pct": _to_float(item.get("阶段涨跌幅")),
            }
        )
    return rows


def _fund_mainline_from_histories(kind: str, existing_errors: list[str] | None = None) -> tuple[list[dict], str, list[str]]:
    base = _sector_funds_payload(kind)
    candidates = [row["name"] for row in (base.get("items") or [])[:20] if row.get("name")]
    rows = []
    errors = list(existing_errors or [])
    for name in candidates[:12]:
        try:
            history = _sector_fund_history(kind, name)
        except Exception as error:
            errors.append(f"{name}: {error}")
            continue
        if not history:
            continue
        recent5 = history[-5:]
        recent3 = history[-3:]
        net3 = round(sum(row["main_net_yi"] for row in recent3), 2)
        net5 = round(sum(row["main_net_yi"] for row in recent5), 2)
        streak = _positive_streak(history)
        if streak >= 3 or net3 > 0:
            rows.append(
                {
                    "name": name,
                    "kind": kind,
                    "net3_yi": net3,
                    "net5_yi": net5,
                    "streak": streak,
                    "latest_date": recent5[-1]["date"],
                    "latest_net_yi": recent5[-1]["main_net_yi"],
                    "history": recent5,
                }
            )
    return rows, "东方财富历史板块资金流 / akshare", errors


def _sector_fund_history(kind: str, name: str) -> list[dict]:
    import akshare as ak

    if kind == "industry":
        df = ak.stock_sector_fund_flow_hist(symbol=name)
    else:
        df = ak.stock_concept_fund_flow_hist(symbol=name)
    columns = list(df.columns)

    def pick(*names):
        for item in names:
            if item in columns:
                return item
        return None

    date_col = pick("日期")
    main_col = pick("主力净流入-净额", "今日主力净流入-净额", "净额")
    pct_col = pick("涨跌幅", "今日涨跌幅")
    rows = []
    for _, item in df.iterrows():
        main_net = _to_float(item.get(main_col)) if main_col else 0
        rows.append(
            {
                "date": str(item.get(date_col) or "")[:10],
                "main_net_yi": round(main_net / 100000000, 2) if abs(main_net) > 10000 else round(main_net, 2),
                "change_pct": _to_float(item.get(pct_col)) if pct_col else None,
            }
        )
    return [row for row in rows if row["date"]]


def _positive_streak(rows: list[dict]) -> int:
    streak = 0
    for row in reversed(rows):
        if row["main_net_yi"] > 0:
            streak += 1
        else:
            break
    return streak


def _fund_mainline_summary(rows: list[dict], label: str, errors: list[str]) -> dict:
    if not rows:
        detail = "阶段资金榜和东方财富历史资金流当前不可用，或候选板块未形成连续净流入。"
        if errors:
            detail = f"{detail} 最近错误示例：{errors[0]}"
        return {"conclusion": f"暂无{label}资金主线。", "analysis": [detail]}
    leader = rows[0]
    continuity = leader.get("continuity_label") or f"连续净流入 {leader['streak']} 日"
    return {
        "conclusion": f"{label}资金主线领跑：{leader['name']}，{continuity}，近 5 日主力净流 {leader['net5_yi']:+.2f} 亿元。",
        "analysis": [
            f"本榜单优先使用同花顺{label}3/5/10日阶段资金榜，失败时再尝试东方财富单板块历史资金流。",
            "阶段榜要求近 3 日与近 5 日主力净流入同时为正；排序优先看为正阶段数，其次比较近 5 日与近 3 日净流入规模。",
            "阶段榜适合识别资金主线，不把即时资金冒充连续历史资金。",
        ],
    }


def _macro_commodities_payload() -> dict:
    bond = _macro_bond_payload()
    commodities = _macro_futures_payload()
    fx = _macro_fx_payload()
    rows = [*bond.get("items", []), *fx.get("items", []), *commodities.get("items", [])]
    available_count = sum(1 for row in rows if row.get("available", True))
    return {
        "available": available_count > 0,
        "source": "FRED 美债/美元/人民币日频 + akshare 中国国债收益率 + 新浪国内期货主连",
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "summary": _macro_summary(rows),
        "groups": {"rates": bond, "fx": fx, "commodities": commodities},
        "items": rows,
    }


def _macro_bond_payload() -> dict:
    items = []
    fred_errors = []
    for series_id in ("DGS10", "DGS2"):
        try:
            items.append(_fred_latest_item(series_id))
        except Exception as error:
            fred_errors.append(f"{series_id}: {error}")
    try:
        import akshare as ak

        df = ak.bond_zh_us_rate(start_date=(datetime.now(TZ).date() - timedelta(days=180)).strftime("%Y%m%d"))
        clean = df.dropna(how="all")
        latest = clean.iloc[-1].to_dict()
        previous = clean.iloc[-2].to_dict() if len(clean) >= 2 else latest
        cn_value = _to_float(latest.get("中国国债收益率10年"))
        cn_prev = _to_float(previous.get("中国国债收益率10年"))
        items.append({"group": "利率", "name": "中国10Y国债", "value": round(cn_value, 4), "change": round(cn_value - cn_prev, 4), "unit": "%", "date": str(latest.get("日期")), "available": True, "source": "akshare/bond_zh_us_rate"})
        us10 = next((row for row in items if row["name"] == "美国10Y国债" and row.get("available") is not False), None)
        if us10:
            spread = round(cn_value - _to_float(us10.get("value")), 4)
            previous_spread = round(cn_prev - (_to_float(us10.get("value")) - _to_float(us10.get("change"))), 4)
            items.append({"group": "利率", "name": "中美10Y利差", "value": spread, "change": round(spread - previous_spread, 4), "unit": "%", "date": str(latest.get("日期")), "available": True, "source": "akshare + FRED"})
    except Exception as error:
        items.append(_unavailable_macro("利率", "中国10Y国债", error, source="akshare/bond_zh_us_rate"))
    for item in fred_errors:
        items.append(_unavailable_macro("利率", item.split(":", 1)[0], item, source="FRED"))
    return {"available": any(item.get("available", True) for item in items), "items": items}


def _fred_latest_item(series_id: str) -> dict:
    meta = FRED_SERIES[series_id]
    response = requests.get(FRED_CSV_URL, params={"id": series_id}, timeout=20)
    response.raise_for_status()
    reader = csv.DictReader(io.StringIO(response.text))
    rows = []
    for row in reader:
        value = row.get(series_id)
        if value and value != ".":
            rows.append({"date": row.get("observation_date", ""), "value": _to_float(value)})
    if not rows:
        raise RuntimeError("FRED 未返回有效观测值")
    latest = rows[-1]
    previous = rows[-2] if len(rows) >= 2 else latest
    return {
        "group": meta["group"],
        "name": meta["name"],
        "value": round(latest["value"], 4),
        "change": round(latest["value"] - previous["value"], 4),
        "unit": meta["unit"],
        "date": latest["date"],
        "available": True,
        "source": meta["source"],
    }


def _macro_fx_payload() -> dict:
    items = []
    for series_id in ("DTWEXBGS", "DEXCHUS"):
        try:
            items.append(_fred_latest_item(series_id))
        except Exception as error:
            items.append(_unavailable_macro("汇率", FRED_SERIES[series_id]["name"], error, source=FRED_SERIES[series_id]["source"]))
    return {"available": any(item.get("available", True) for item in items), "items": items}


def _macro_futures_payload() -> dict:
    start_date = (datetime.now(TZ).date() - timedelta(days=30)).strftime("%Y%m%d")
    end_date = datetime.now(TZ).date().strftime("%Y%m%d")
    items = []
    try:
        import akshare as ak

        for contract in FUTURES_MAIN_CONTRACTS:
            try:
                df = ak.futures_main_sina(symbol=contract["symbol"], start_date=start_date, end_date=end_date)
                rows = df.dropna(how="all")
                if rows.empty:
                    raise RuntimeError("未返回主连行情")
                latest = rows.iloc[-1].to_dict()
                previous = rows.iloc[-2].to_dict() if len(rows) >= 2 else latest
                value = _to_float(latest.get("收盘价"))
                prev = _to_float(previous.get("收盘价"))
                change_pct = ((value / prev - 1) * 100) if prev else 0
                items.append(
                    {
                        "group": contract["group"],
                        "name": contract["name"],
                        "value": round(value, 4),
                        "change": round(change_pct, 2),
                        "unit": contract["unit"],
                        "date": str(latest.get("日期")),
                        "available": True,
                        "source": f"新浪期货主连/{contract['symbol']} / akshare",
                    }
                )
            except Exception as error:
                items.append(_unavailable_macro(contract["group"], contract["name"], error, source=f"新浪期货主连/{contract['symbol']}"))
    except Exception as error:
        items.append(_unavailable_macro("商品", "热门商品", error, source="akshare/futures_main_sina"))
    return {"available": any(item.get("available", True) for item in items), "items": items}


def _find_row_by_keywords(df, keywords: tuple[str, ...]) -> dict | None:
    for _, item in df.iterrows():
        text = " ".join(str(value) for value in item.to_dict().values())
        if any(keyword in text for keyword in keywords):
            return item.to_dict()
    return None


def _unavailable_macro(group: str, name: str, error, source: str | None = None) -> dict:
    return {"group": group, "name": name, "value": None, "change": None, "unit": "", "date": "-", "available": False, "detail": str(error), "source": source or ""}


def _macro_summary(rows: list[dict]) -> dict:
    available = [row for row in rows if row.get("available", True)]
    risk_notes = []
    us10 = next((row for row in available if row["name"] == "美国10Y国债"), None)
    cny = next((row for row in available if "人民币" in row["name"]), None)
    oil = next((row for row in available if row["name"] == "原油"), None)
    copper = next((row for row in available if row["name"] == "铜"), None)
    if us10 and us10.get("change", 0) > 0:
        risk_notes.append("美债收益率上行，对全球权益估值偏压制。")
    if cny and cny.get("change", 0) > 0:
        risk_notes.append("人民币汇率走弱时，需关注外资与大盘权重承压。")
    if oil and oil.get("change", 0) > 0:
        risk_notes.append("原油上涨可能推升通胀预期，对航空、化工成本端有扰动。")
    if copper and copper.get("change", 0) > 0:
        risk_notes.append("铜价走强通常对应全球工业需求或再通胀交易升温。")
    return {
        "conclusion": f"宏观商品当前可用指标 {len(available)} 个。",
        "analysis": risk_notes or ["当前宏观商品指标未形成单边压力，重点继续观察美元、美债、人民币与工业品共振。"],
    }


def _parse_mcp_sse(text: str) -> dict:
    for line in text.splitlines():
        if line.startswith("data:"):
            return json.loads(line.removeprefix("data:").strip())
    return json.loads(text)


def _json_safe_row(row: dict) -> dict:
    output = {}
    for key, value in row.items():
        if isinstance(value, (date, datetime)):
            output[key] = value.isoformat()
        elif isinstance(value, Decimal):
            output[key] = float(value)
        else:
            output[key] = value
    return output


def _calendar_days(start, end) -> int:
    if not start or not end:
        return 0
    if isinstance(start, str):
        start = date.fromisoformat(start[:10])
    if isinstance(end, str):
        end = date.fromisoformat(end[:10])
    return max(1, (end - start).days + 1)


def _coverage_status(row: dict) -> str:
    days = _calendar_days(row.get("min_date"), row.get("max_date"))
    missing_change = int(row.get("missing_change_pct") or 0)
    missing_close = int(row.get("missing_close") or 0)
    if row["market"] == "A股个股" and days < 180:
        return "历史偏短"
    if missing_close > 0 or missing_change > int(row.get("code_count") or 0):
        return "字段缺失"
    if row["code_count"] <= 3:
        return "样本较少"
    return "可用"


def _data_center_suggestions(markets: list[dict], a_stock: dict | None) -> list[str]:
    rows = []
    if a_stock:
        days = _calendar_days(a_stock.get("min_date"), a_stock.get("max_date"))
        if days < 250:
            rows.append(f"A 股个股当前覆盖约 {days} 个自然日，建议用 RssCast 慢速补齐最近一年，才能支撑趋势分位和长周期资金强弱。")
        if a_stock.get("code_count", 0) < 800:
            rows.append(f"A 股个股池当前 {a_stock.get('code_count')} 只，更适合做核心样本观察；若要全市场热度，需要扩展股票池。")
    if any(row["market"] == "美股波动率" for row in markets):
        rows.append("已存在美股波动率数据，可以纳入 A 股开盘前外围风险判断。")
    if any(row["market"] == "港股指数" for row in markets) and any(row["market"] == "美股指数" for row in markets):
        rows.append("已具备港股、美股指数数据，适合做跨市场风险面板和隔夜风险提示。")
    return rows or ["当前库可用于基础行情展示，建议继续补充资金流、宏观和商品数据。"]


def _risk_score_from_cross_market(latest: list[dict], histories: dict[str, list[dict]]) -> dict:
    def row_for(*keywords):
        for row in latest:
            text = f"{row.get('code', '')} {row.get('name', '')}"
            if any(keyword in text for keyword in keywords):
                return row
        return None

    hsi = row_for("HSI", "恒生")
    nasdaq = row_for("IXIC", "纳斯达克")
    spx = row_for("GSPC", "标普")
    dow = row_for("DJI", "道琼")
    vix = row_for("VIX", "波动率")
    offshore = [row for row in (hsi, nasdaq, spx, dow) if row]
    avg_offshore = _average([float(row.get("change_pct") or 0) for row in offshore]) or 0
    vix_pct = float((vix or {}).get("change_pct") or 0)
    score = _clamp(50 - avg_offshore * 8 + max(vix_pct, 0) * 2.2 - min(vix_pct, 0) * 0.8)
    if score >= 70:
        state = "外围高压"
    elif score >= 58:
        state = "外围偏谨慎"
    elif score <= 38:
        state = "外围风险缓和"
    else:
        state = "外围中性"
    signals = []
    if offshore:
        signals.append(f"港美主要指数平均涨跌幅 {avg_offshore:+.2f}%。")
    if vix:
        signals.append(f"VIX/波动率代理涨跌幅 {vix_pct:+.2f}%。")
    if hsi:
        signals.append(f"港股风险偏好参考：{hsi['name']} {float(hsi.get('change_pct') or 0):+.2f}%。")
    return {"score": round(score), "state": state, "signals": signals, "tracked": len(latest)}


def _cross_market_summary(risk: dict, latest: list[dict]) -> dict:
    conclusion = f"外围风险状态：{risk['state']}，风险分 {risk['score']}。"
    analysis = list(risk.get("signals") or [])
    if not analysis:
        analysis.append("本地库暂未提供足够的港股、美股或波动率数据。")
    if risk["score"] >= 58:
        analysis.append("若 A 股自身情绪分数同步走弱，开盘后需优先观察权重指数承接与北向/主力资金是否继续流出。")
    else:
        analysis.append("外围环境未形成明显压力，A 股开盘后的主线更应看行业资金和市场宽度能否扩散。")
    return {"conclusion": conclusion, "analysis": analysis, "count": len(latest)}


def _normalize_sector_fund_rows(df, label: str) -> list[dict]:
    rows = []
    columns = list(df.columns)

    def pick(*names):
        for name in names:
            if name in columns:
                return name
        return None

    name_col = pick("名称", "板块名称", f"{label}名称", "行业")
    code_col = pick("代码", "板块代码")
    main_col = pick("今日主力净流入-净额", "主力净流入-净额", "主力净流入净额", "净额", "今日主力净流入")
    pct_col = pick("今日涨跌幅", "涨跌幅", "板块涨跌幅", "行业-涨跌幅")
    main_pct_col = pick("今日主力净流入-净占比", "主力净流入-净占比", "主力净流入净占比", "净占比")
    for _, item in df.iterrows():
        main_net = _to_float(item.get(main_col)) if main_col else 0
        change_pct = _to_float(item.get(pct_col)) if pct_col else 0
        main_ratio = _to_float(item.get(main_pct_col)) if main_pct_col else None
        rows.append(
            {
                "code": str(item.get(code_col) or ""),
                "name": str(item.get(name_col) or ""),
                "main_net_yi": round(main_net / 100000000, 2) if abs(main_net) > 10000 else round(main_net, 2),
                "main_net_raw": main_net,
                "main_ratio": main_ratio,
                "change_pct": change_pct,
                "heat": round(abs(main_net / 100000000) if abs(main_net) > 10000 else abs(main_net), 2),
            }
        )
    return sorted(rows, key=lambda row: abs(row["main_net_yi"]), reverse=True)


def _sector_funds_fallback(kind: str, label: str, original_error: str) -> tuple[list[dict], str]:
    fallback_kind = "industry" if kind == "region" else kind
    fallback_label = "行业" if kind == "region" else label
    if fallback_kind not in ("industry", "concept"):
        return [], "东方财富板块资金流 / akshare"
    try:
        import akshare as ak

        if fallback_kind == "industry":
            df = ak.stock_fund_flow_industry(symbol="即时")
            source = "同花顺行业资金流 / akshare" if kind != "region" else "地域资金不可用，降级显示同花顺行业资金流 / akshare"
        else:
            df = ak.stock_fund_flow_concept(symbol="即时")
            source = "同花顺概念资金流 / akshare"
        rows = []
        for _, item in df.iterrows():
            main_net = _to_float(item.get("净额"))
            name = str(item.get("行业") or item.get("概念") or item.get("名称") or item.get("板块") or "")
            if not name:
                continue
            rows.append(
                {
                    "code": "",
                    "name": name,
                    "main_net_yi": round(main_net, 2),
                    "main_net_raw": main_net,
                    "main_ratio": None,
                    "change_pct": _to_float(item.get("行业-涨跌幅") or item.get("概念-涨跌幅") or item.get("涨跌幅")),
                    "heat": abs(round(main_net, 2)),
                    "leader": str(item.get("领涨股") or item.get("龙头股") or ""),
                }
            )
        return sorted(rows, key=lambda row: abs(row["main_net_yi"]), reverse=True), source
    except Exception:
        return [], f"东方财富板块资金流 / akshare；备用源也不可用；原始错误：{original_error}"


def _sector_funds_summary(rows: list[dict], label: str) -> dict:
    if not rows:
        return {"conclusion": f"暂无{label}主力资金数据。", "analysis": []}
    inflow = [row for row in rows if row["main_net_yi"] > 0]
    outflow = [row for row in rows if row["main_net_yi"] < 0]
    strongest = max(rows, key=lambda row: row["main_net_yi"])
    weakest = min(rows, key=lambda row: row["main_net_yi"])
    net_total = round(sum(row["main_net_yi"] for row in rows), 2)
    conclusion = f"{label}资金净额合计 {net_total:+.0f} 亿元，净流入板块 {len(inflow)} 个，净流出板块 {len(outflow)} 个。"
    return {
        "conclusion": conclusion,
        "strongest": strongest,
        "weakest": weakest,
        "analysis": [
            f"主力净流入最强：{strongest['name']} {strongest['main_net_yi']:+.2f} 亿元。",
            f"主力净流出最强：{weakest['name']} {weakest['main_net_yi']:+.2f} 亿元。",
            "热力图面积按主力净额绝对值放大，颜色区分净流入与净流出，用于观察资金集中度而不是只看涨跌幅。",
        ],
    }


def _to_float(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, str):
        value = value.replace(",", "").replace("%", "").strip()
        if value in ("", "-", "--"):
            return 0.0
    try:
        number = float(value)
        return number if math.isfinite(number) else 0.0
    except Exception:
        return 0.0


def _enrich_quote(row: dict, name: str) -> dict:
    close = row.get("last_price", row.get("close"))
    pre_close = row.get("pre_close")
    change_pct = row.get("change_pct")
    if change_pct is None and close is not None and pre_close:
        change_pct = (float(close) - float(pre_close)) / float(pre_close)
    return {
        **row,
        "name": name,
        "close": close,
        "change_pct_display": round(float(change_pct or 0) * 100, 2),
        "amount_yi": round(float(row.get("amount") or 0) / 100000000, 2),
        "volume_wan": round(float(row.get("volume") or 0) / 10000, 2),
    }


def _group_history(rows: list[dict], meta: dict[str, dict]) -> dict:
    grouped: dict[str, list[dict]] = {}
    for row in rows:
        grouped.setdefault(row["code"], []).append(row)
    return {
        code: {
            "name": meta.get(code, {}).get("name", code),
            "rows": _history_metrics(sorted(items, key=lambda row: row.get("timeString", ""))),
        }
        for code, items in grouped.items()
    }


def _history_metrics(rows: list[dict]) -> list[dict]:
    closes = [float(row["close"]) for row in rows if row.get("close") is not None]
    output = []
    for index, row in enumerate(rows):
        close = float(row["close"])
        ma5 = _moving_average(closes, index, 5)
        ma20 = _moving_average(closes, index, 20)
        pct5 = _period_change(closes, index, 5)
        pct20 = _period_change(closes, index, 20)
        output.append(
            {
                **row,
                "date": str(row.get("timeString", ""))[:10],
                "ma5": ma5,
                "ma20": ma20,
                "pct5": pct5,
                "pct20": pct20,
                "amount_yi": round(float(row.get("amount") or 0) / 100000000, 2),
            }
        )
    return output


def _moving_average(values: list[float], index: int, window: int) -> float | None:
    if index + 1 < window:
        return None
    scope = values[index + 1 - window : index + 1]
    return round(sum(scope) / len(scope), 2)


def _period_change(values: list[float], index: int, window: int) -> float | None:
    if index < window or values[index - window] == 0:
        return None
    return round((values[index] / values[index - window] - 1) * 100, 2)


def _latest_time(rows: list[dict]) -> str | None:
    times = [row.get("timeString") for row in rows if row.get("timeString")]
    return max(times) if times else None


def _index_summary(rows: list[dict], histories: dict) -> dict:
    if not rows:
        return {"conclusion": "暂无 A 股指数数据。", "analysis": []}
    up_count = sum(1 for row in rows if float(row.get("change_pct") or 0) > 0)
    down_count = sum(1 for row in rows if float(row.get("change_pct") or 0) < 0)
    strongest = max(rows, key=lambda row: float(row.get("change_pct") or 0))
    weakest = min(rows, key=lambda row: float(row.get("change_pct") or 0))
    sh_amount = next((row["amount_yi"] for row in rows if row["code"] == "000001"), 0)
    sz_amount = next((row["amount_yi"] for row in rows if row["code"] == "399001"), 0)
    total_amount = round(sh_amount + sz_amount, 2)
    if up_count >= 4:
        conclusion = "核心指数多数上涨，A 股风险偏好偏积极。"
    elif down_count >= 4:
        conclusion = "核心指数多数回调，A 股风险偏好偏谨慎。"
    else:
        conclusion = "核心指数表现分化，A 股处于结构性震荡。"
    return {
        "conclusion": conclusion,
        "up_count": up_count,
        "down_count": down_count,
        "strongest": strongest,
        "weakest": weakest,
        "core_amount_yi": total_amount,
        "analysis": [
            f"核心指数中 {up_count} 个上涨、{down_count} 个下跌，最强为 {strongest['name']}，最弱为 {weakest['name']}。",
            f"上证指数与深证成指成交额合计约 {total_amount:.0f} 亿元，用于观察市场量能变化。",
            _style_signal(histories),
        ],
    }


def _style_signal(histories: dict) -> str:
    growth = _latest_history_change(histories, "399006")
    bluechip = _latest_history_change(histories, "000300")
    if growth is None or bluechip is None:
        return "风格信号暂不完整，需等待更多历史数据。"
    if growth > bluechip:
        return f"近 20 日创业板相对沪深300更强，成长风格阶段占优约 {growth - bluechip:.2f} 个百分点。"
    return f"近 20 日沪深300相对创业板更强，大盘蓝筹阶段占优约 {bluechip - growth:.2f} 个百分点。"


def _latest_history_change(histories: dict, code: str) -> float | None:
    rows = histories.get(code, {}).get("rows") or []
    return rows[-1].get("pct20") if rows else None


def _load_watchlist() -> list[dict]:
    if not WATCHLIST_PATH.exists():
        _save_watchlist(DEFAULT_WATCHLIST)
        return DEFAULT_WATCHLIST.copy()
    with WATCHLIST_PATH.open(encoding="utf-8") as file:
        items = json.load(file)
    changed = False
    for item in items:
        if "category" not in item or not item.get("category"):
            item["category"] = DEFAULT_WATCHLIST_CATEGORY
            changed = True
        if "name" not in item or not item.get("name"):
            item["name"] = item.get("code", "")
            changed = True
    if changed:
        _save_watchlist(items)
    return items


def _save_watchlist(items: list[dict]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with WATCHLIST_PATH.open("w", encoding="utf-8") as file:
        json.dump(items, file, ensure_ascii=False, indent=2)


def _watchlist_payload(items: list[dict]) -> dict:
    codes = [item["code"] for item in items]
    quotes = _rsscast_call("StockPriceQuery", {"codes": codes}) if codes else []
    item_map = {item["code"]: item for item in items}
    local_names = _market_data_name_map(codes)
    changed = False
    for item in items:
        local_name = local_names.get(item["code"])
        if local_name and (not item.get("name") or item.get("name") == item["code"]):
            item["name"] = local_name
            changed = True
    if changed:
        _save_watchlist(items)
    rows = []
    for row in quotes:
        item = item_map.get(row["code"], {})
        saved_name = item.get("name")
        display_name = (
            saved_name
            if saved_name and saved_name != row["code"]
            else local_names.get(row["code"]) or row.get("name") or row["code"]
        )
        enriched = _enrich_quote(row, display_name)
        enriched["category"] = item.get("category") or DEFAULT_WATCHLIST_CATEGORY
        rows.append(enriched)
    rows = sorted(rows, key=lambda row: codes.index(row["code"]) if row["code"] in codes else 999)
    categories = _watchlist_categories(items, rows)
    return {
        "source": "RssCast MCP",
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "trade_time": _latest_time(rows),
        "items": rows,
        "categories": categories,
        "summary": _watchlist_summary(rows),
    }


def _watchlist_categories(items: list[dict], rows: list[dict]) -> list[dict]:
    counts: dict[str, int] = {}
    for item in items:
        category = item.get("category") or DEFAULT_WATCHLIST_CATEGORY
        counts[category] = counts.get(category, 0) + 1
    for row in rows:
        category = row.get("category") or DEFAULT_WATCHLIST_CATEGORY
        counts.setdefault(category, 0)
    ordered = [{"name": "全部", "count": len(items)}]
    ordered.extend({"name": name, "count": count} for name, count in sorted(counts.items()))
    return ordered


def _market_data_name_map(codes: list[str]) -> dict[str, str]:
    if not codes:
        return {}
    try:
        import pymysql

        connection = pymysql.connect(host="127.0.0.1", user="root", database=MARKET_DATA_DB, charset="utf8mb4")
    except Exception:
        return {}
    placeholders = ",".join(["%s"] * len(codes))
    sql = f"""
        SELECT REPLACE(REPLACE(code, '.SS', ''), '.SZ', '') AS plain_code, MAX(name) AS name
        FROM daily_data
        WHERE REPLACE(REPLACE(code, '.SS', ''), '.SZ', '') IN ({placeholders})
          AND name IS NOT NULL
          AND name <> ''
        GROUP BY plain_code
    """
    try:
        with connection.cursor() as cursor:
            cursor.execute(sql, codes)
            return {str(code): str(name) for code, name in cursor.fetchall() if name}
    except Exception:
        return {}
    finally:
        connection.close()


def _watchlist_name(code: str) -> str:
    for item in _load_watchlist():
        if item["code"] == code:
            return item.get("name") or code
    return code


def _watchlist_summary(rows: list[dict]) -> dict:
    if not rows:
        return {"conclusion": "自选池为空。", "up_count": 0, "down_count": 0}
    up_count = sum(1 for row in rows if float(row.get("change_pct") or 0) > 0)
    down_count = sum(1 for row in rows if float(row.get("change_pct") or 0) < 0)
    strongest = max(rows, key=lambda row: float(row.get("change_pct") or 0))
    weakest = min(rows, key=lambda row: float(row.get("change_pct") or 0))
    conclusion = f"自选池 {len(rows)} 只股票中 {up_count} 只上涨、{down_count} 只下跌。"
    return {"conclusion": conclusion, "up_count": up_count, "down_count": down_count, "strongest": strongest, "weakest": weakest}


def _stock_signals(quote: dict | None, history: list[dict]) -> list[str]:
    if not history:
        return ["历史 K 线不足，暂无法生成技术信号。"]
    latest = history[-1]
    signals = []
    if quote:
        signals.append(f"最新收盘/成交价 {quote['close']}，涨跌幅 {quote['change_pct_display']:+.2f}%。")
    if latest.get("ma20") is not None:
        relation = "站上" if float(latest["close"]) >= float(latest["ma20"]) else "跌破"
        signals.append(f"价格{relation} MA20，MA20 当前约 {latest['ma20']}。")
    if latest.get("pct20") is not None:
        signals.append(f"近 20 个交易日涨跌幅 {latest['pct20']:+.2f}%。")
    return signals


def _sentiment_payload(indices: list[dict], histories: dict, breadth_stats: dict | None, watch_payload: dict) -> dict:
    components = [
        _index_emotion_component(indices),
        _breadth_emotion_component(breadth_stats),
        _volume_emotion_component(indices, histories),
        _style_emotion_component(indices, histories),
        _watchlist_emotion_component(watch_payload),
    ]
    score = round(sum(item["score"] * item["weight"] for item in components) / sum(item["weight"] for item in components))
    label = _sentiment_label(score)
    tag = _sentiment_tag(score, components)
    trade_time = _latest_time(indices) or watch_payload.get("trade_time")
    conclusion = _sentiment_conclusion(score, tag, components)
    analysis = _sentiment_analysis(components, breadth_stats, watch_payload)
    return {
        "source": "RssCast MCP + market-breadth-heatmap skill",
        "generated_at": datetime.now(TZ).isoformat(timespec="seconds"),
        "trade_time": trade_time,
        "score": score,
        "label": label,
        "tag": tag,
        "conclusion": conclusion,
        "components": components,
        "analysis": analysis,
        "evidence": {
            "indices": indices,
            "breadth": breadth_stats,
            "watchlist": watch_payload,
        },
    }


def _index_emotion_component(indices: list[dict]) -> dict:
    if not indices:
        return _component("index", "指数动能", 50, "数据不足", "未取得核心指数行情。", 0.24)
    pct_values = [float(row.get("change_pct_display") or 0) for row in indices]
    up_ratio = sum(1 for value in pct_values if value > 0) / len(pct_values)
    avg_pct = sum(pct_values) / len(pct_values)
    score = _clamp(50 + avg_pct * 11 + (up_ratio - 0.5) * 34)
    strongest = max(indices, key=lambda row: float(row.get("change_pct_display") or 0))
    weakest = min(indices, key=lambda row: float(row.get("change_pct_display") or 0))
    if up_ratio >= 0.67:
        state = "指数共振上行"
    elif up_ratio <= 0.33:
        state = "指数同步回调"
    else:
        state = "指数结构分化"
    detail = f"{len(indices)} 个核心指数中 {round(up_ratio * len(indices))} 个上涨，均值 {avg_pct:+.2f}%；最强 {strongest['name']}，最弱 {weakest['name']}。"
    return _component("index", "指数动能", score, state, detail, 0.24)


def _breadth_emotion_component(stats: dict | None) -> dict:
    if not stats:
        return _component("breadth", "市场宽度", 50, "宽度缺失", "未读取到市场宽度聚合数据，暂按中性处理。", 0.28)
    latest_values = stats.get("latest_values") or []
    strong_count = sum(1 for item in latest_values if item["value"] >= 60)
    weak_count = sum(1 for item in latest_values if item["value"] < 30)
    improved = stats.get("improved") or []
    weakened = stats.get("weakened") or []
    improvement = (improved[0]["change"] if improved else 0) + (weakened[0]["change"] if weakened else 0)
    score = _clamp(float(stats["average"]) + improvement * 0.15 + (strong_count - weak_count) * 1.2)
    if score >= 60:
        state = "赚钱效应扩散"
    elif score >= 45:
        state = "宽度中性分化"
    elif score >= 30:
        state = "宽度收缩"
    else:
        state = "赚钱效应偏冷"
    detail = f"行业 MA20 站上率均值 {stats['average']:.1f}%，强势行业 {strong_count} 个，弱势行业 {weak_count} 个。"
    return _component("breadth", "市场宽度", score, state, detail, 0.28)


def _volume_emotion_component(indices: list[dict], histories: dict) -> dict:
    sh = next((row for row in indices if row["code"] == "000001"), None)
    sz = next((row for row in indices if row["code"] == "399001"), None)
    latest_amount = float((sh or {}).get("amount_yi") or 0) + float((sz or {}).get("amount_yi") or 0)
    history_amounts = []
    for code in ("000001", "399001"):
      rows = histories.get(code, {}).get("rows") or []
      history_amounts.append([float(row.get("amount_yi") or 0) for row in rows[-20:]])
    paired = [sum(values) for values in zip(*history_amounts)] if all(history_amounts) else []
    avg_amount = sum(paired) / len(paired) if paired else latest_amount
    amount_ratio = latest_amount / avg_amount if avg_amount else 1
    avg_index_pct = sum(float(row.get("change_pct_display") or 0) for row in indices) / len(indices) if indices else 0
    directional_adjustment = 10 if avg_index_pct > 0 else -10 if avg_index_pct < -0.5 else 0
    score = _clamp(50 + (amount_ratio - 1) * 65 + directional_adjustment)
    if avg_index_pct < -0.5 and amount_ratio >= 1.03:
        state = "放量降风险"
        score = min(score, 38)
    elif avg_index_pct > 0.3 and amount_ratio >= 1.03:
        state = "放量承接"
    elif amount_ratio < 0.92:
        state = "缩量观望"
    else:
        state = "量能中性"
    detail = f"上证+深成指成交额约 {latest_amount:.0f} 亿元，约为近 20 日均值的 {amount_ratio:.2f} 倍。"
    return _component("volume", "成交量能", score, state, detail, 0.18)


def _style_emotion_component(indices: list[dict], histories: dict) -> dict:
    pct_today = {row["code"]: float(row.get("change_pct_display") or 0) for row in indices}
    high_beta_today = _average([pct_today.get("399006"), pct_today.get("000688"), pct_today.get("000905")])
    core_today = _average([pct_today.get("000300"), pct_today.get("000001")])
    high_beta_20 = _average([_latest_history_change(histories, code) for code in ("399006", "000688", "000905")])
    core_20 = _average([_latest_history_change(histories, code) for code in ("000300", "000001")])
    spread_today = (high_beta_today or 0) - (core_today or 0)
    spread_20 = (high_beta_20 or 0) - (core_20 or 0)
    score = _clamp(50 + spread_today * 12 + spread_20 * 0.9 + (high_beta_today or 0) * 6)
    if high_beta_today is not None and high_beta_today < -1.0:
        score = min(score, 42)
    if spread_today < -0.35:
        score = min(score, 45)
    if spread_today > 0.35 and spread_20 > 0 and (high_beta_today or 0) > 0:
        state = "高 beta 占优"
    elif spread_today < -0.35 or (high_beta_today is not None and high_beta_today < -1.0):
        state = "高 beta 降温"
    elif core_today is not None and high_beta_today is not None and core_today > high_beta_today:
        state = "权重相对抗跌"
    else:
        state = "风格均衡"
    detail = f"高 beta 当日相对核心指数 {spread_today:+.2f} 个百分点，近 20 日相对强弱 {spread_20:+.2f} 个百分点。"
    return _component("style", "风格风险偏好", score, state, detail, 0.18)


def _watchlist_emotion_component(watch_payload: dict) -> dict:
    rows = watch_payload.get("items") or []
    if not rows:
        return _component("watchlist", "自选股状态", 50, "自选池为空", "暂无自选股，暂按中性处理。", 0.12)
    pct_values = [float(row.get("change_pct_display") or 0) for row in rows]
    up_ratio = sum(1 for value in pct_values if value > 0) / len(pct_values)
    avg_pct = sum(pct_values) / len(pct_values)
    score = _clamp(50 + avg_pct * 10 + (up_ratio - 0.5) * 30)
    if up_ratio >= 0.67:
        state = "关注池偏强"
    elif up_ratio <= 0.33:
        state = "关注池偏弱"
    else:
        state = "关注池分化"
    summary = watch_payload.get("summary") or {}
    strongest = summary.get("strongest", {})
    weakest = summary.get("weakest", {})
    detail = f"自选池 {len(rows)} 只股票中 {round(up_ratio * len(rows))} 只上涨，均值 {avg_pct:+.2f}%；最强 {strongest.get('name', '-')}，最弱 {weakest.get('name', '-')}。"
    return _component("watchlist", "自选股状态", score, state, detail, 0.12)


def _component(key: str, name: str, score: float, state: str, detail: str, weight: float) -> dict:
    return {"key": key, "name": name, "score": round(_clamp(score)), "state": state, "detail": detail, "weight": weight}


def _sentiment_label(score: int) -> str:
    if score >= 75:
        return "过热"
    if score >= 60:
        return "偏热"
    if score >= 45:
        return "中性"
    if score >= 30:
        return "谨慎"
    return "偏冷"


def _sentiment_tag(score: int, components: list[dict]) -> str:
    states = {item["key"]: item["state"] for item in components}
    if score >= 75:
        return "全面 Risk-on"
    if score >= 60:
        return "结构性 Risk-on" if states.get("style") in ("高 beta 占优", "风格均衡") else "权重护盘"
    if score >= 45:
        return "分化轮动"
    if states.get("volume") == "放量降风险":
        return "放量杀跌"
    if states.get("style") == "高 beta 降温":
        return "高 beta 降温"
    return "缩量观望"


def _sentiment_conclusion(score: int, tag: str, components: list[dict]) -> str:
    weakest = min(components, key=lambda item: item["score"])
    strongest = max(components, key=lambda item: item["score"])
    if score >= 60:
        return f"今日 A 股情绪处于{_sentiment_label(score)}区间，交易状态更接近「{tag}」；主要支撑来自{strongest['name']}，但仍需观察{weakest['name']}是否拖累扩散。"
    if score >= 45:
        return f"今日 A 股情绪处于中性震荡区间，交易状态更接近「{tag}」；市场不是单边风险偏好，核心矛盾在{strongest['name']}与{weakest['name']}之间的拉扯。"
    return f"今日 A 股情绪偏谨慎，交易状态更接近「{tag}」；短线需要先看{weakest['name']}能否修复，否则风险偏好仍偏收缩。"


def _sentiment_analysis(components: list[dict], breadth_stats: dict | None, watch_payload: dict) -> list[str]:
    rows = [f"{item['name']}：{item['state']}，评分 {item['score']}。{item['detail']}" for item in components]
    if breadth_stats:
        rows.append(
            f"行业层面最强为 {breadth_stats['strongest']['category']}，最弱为 {breadth_stats['weakest']['category']}；如果指数继续回调但宽度不再恶化，行情更可能是结构切换而不是全面 risk-off。"
        )
    summary = watch_payload.get("summary") or {}
    if summary.get("strongest") and summary.get("weakest"):
        rows.append(f"自选池中最强为 {summary['strongest']['name']}，最弱为 {summary['weakest']['name']}，用于观察关注资产是否领先或滞后大盘情绪。")
    return rows


def _sentiment_p1_payload(payload: dict, breadth_stats: dict | None) -> dict:
    history = _sentiment_history_payload(payload)
    trade_date = str(payload.get("trade_time") or payload["generated_at"])[:10]
    return {
        **history,
        "limit_mood": _limit_mood_payload(trade_date),
        "capital_flow": _capital_flow_payload(),
        "northbound": _northbound_payload(),
        "industry_crowding": _industry_crowding_payload(breadth_stats),
        "extreme": _sentiment_extreme_payload(history),
    }


def _limit_mood_payload(trade_date: str) -> dict:
    date_key = trade_date.replace("-", "")
    up_payload = _eastmoney_get(EASTMONEY_LIMIT_UP_URL, {"d": date_key, "Pageindex": 0, "pagesize": 10, "sort": "fbt:asc"})
    down_payload = _eastmoney_get(EASTMONEY_LIMIT_DOWN_URL, {"d": date_key, "Pageindex": 0, "pagesize": 10, "sort": "fbt:asc"})
    up_data = up_payload.get("data") if up_payload else None
    down_data = down_payload.get("data") if down_payload else None
    if not isinstance(up_data, dict) or not isinstance(down_data, dict):
        return {
            "available": False,
            "source": "东方财富涨跌停池",
            "date": trade_date,
            "score": None,
            "state": "数据暂不可用",
            "detail": "当前公开接口未返回有效涨跌停池数据，暂不纳入情绪总分，避免误判。",
            "limit_up_count": None,
            "limit_down_count": None,
            "samples": [],
        }
    up_pool = up_data.get("pool") or []
    down_pool = down_data.get("pool") or []
    up_count = int(up_data.get("tc") or len(up_pool))
    down_count = int(down_data.get("tc") or len(down_pool))
    score = _clamp(50 + (up_count - down_count) * 1.2)
    if up_count >= down_count * 2 and up_count >= 30:
        state = "涨停扩散"
    elif down_count >= up_count:
        state = "跌停压力"
    else:
        state = "涨跌停中性"
    return {
        "available": True,
        "source": "东方财富涨跌停池",
        "date": trade_date,
        "score": round(score),
        "state": state,
        "detail": f"涨停 {up_count} 只，跌停 {down_count} 只，涨跌停差额 {up_count - down_count:+d}。",
        "limit_up_count": up_count,
        "limit_down_count": down_count,
        "samples": [{"code": item.get("c"), "name": item.get("n")} for item in up_pool[:5]],
    }


def _capital_flow_payload() -> dict:
    params = {
        "fltt": 2,
        "invt": 2,
        "fields": "f1,f2,f3,f4,f12,f13,f14,f62,f66,f69,f72,f75,f78,f81,f84,f87,f124",
        "secids": "1.000001,0.399001",
    }
    payload = _eastmoney_get(EASTMONEY_INDEX_FLOW_URL, params)
    diff = ((payload or {}).get("data") or {}).get("diff") or []
    rows = []
    for item in diff:
        main_net = float(item.get("f62") or 0) / 100000000
        large_net = float(item.get("f66") or 0) / 100000000
        rows.append(
            {
                "code": item.get("f12"),
                "name": item.get("f14"),
                "change_pct": item.get("f3"),
                "main_net_yi": round(main_net, 2),
                "large_net_yi": round(large_net, 2),
                "time": datetime.fromtimestamp(int(item.get("f124") or 0), TZ).isoformat(timespec="seconds") if item.get("f124") else None,
            }
        )
    total_main = round(sum(row["main_net_yi"] for row in rows), 2)
    if not rows:
        return {"available": False, "source": "东方财富指数资金流", "state": "数据暂不可用", "detail": "未取得指数资金流。", "rows": []}
    state = "主力净流入" if total_main > 0 else "主力净流出" if total_main < 0 else "资金均衡"
    score = _clamp(50 + total_main / 80)
    return {
        "available": True,
        "source": "东方财富指数资金流",
        "score": round(score),
        "state": state,
        "detail": f"上证指数与深证成指主力净流合计 {total_main:+.0f} 亿元。",
        "total_main_net_yi": total_main,
        "rows": rows,
    }


def _northbound_payload() -> dict:
    payload = _eastmoney_get(
        EASTMONEY_KAMT_URL,
        {"fields1": "f1,f2,f3,f4", "fields2": "f51,f52,f53,f54,f55,f56,f57,f58,f59,f60,f61,f62,f63"},
    )
    data = (payload or {}).get("data") or {}
    channels = []
    for key, name in (("hk2sh", "沪股通"), ("hk2sz", "深股通"), ("sh2hk", "港股通沪"), ("sz2hk", "港股通深")):
        row = data.get(key) or {}
        if row:
            channels.append(
                {
                    "key": key,
                    "name": name,
                    "date": row.get("date2") or row.get("date"),
                    "net_buy_yi": round(float(row.get("netBuyAmt") or 0) / 10000, 2),
                    "buy_sell_yi": round(float(row.get("buySellAmt") or 0) / 10000, 2),
                    "status": row.get("status"),
                }
            )
    north = [row for row in channels if row["key"] in ("hk2sh", "hk2sz")]
    south = [row for row in channels if row["key"] in ("sh2hk", "sz2hk")]
    north_net = round(sum(row["net_buy_yi"] for row in north), 2)
    south_net = round(sum(row["net_buy_yi"] for row in south), 2)
    if not channels:
        return {"available": False, "source": "东方财富沪深港通资金", "state": "数据暂不可用", "detail": "未取得沪深港通资金数据。", "channels": []}
    state = "北向净流入" if north_net > 0 else "北向净流出" if north_net < 0 else "北向无净流记录"
    return {
        "available": True,
        "source": "东方财富沪深港通资金",
        "state": state,
        "detail": f"北向资金净买入 {north_net:+.2f} 亿元，南向资金净买入 {south_net:+.2f} 亿元。",
        "north_net_yi": north_net,
        "south_net_yi": south_net,
        "channels": channels,
    }


def _industry_crowding_payload(stats: dict | None) -> dict:
    if not stats:
        return {"available": False, "source": "market-breadth aggregated_data.json", "state": "数据暂不可用", "detail": "未读取到行业宽度数据。", "rows": []}
    latest_values = stats.get("latest_values") or []
    rows = []
    for item in latest_values:
        value = float(item["value"])
        if value >= 75:
            level = "过热"
        elif value >= 60:
            level = "偏热"
        elif value < 20:
            level = "冰点"
        elif value < 30:
            level = "偏冷"
        else:
            level = "中性"
        rows.append({"category": item["category"], "value": value, "level": level})
    hot = [row for row in rows if row["value"] >= 60]
    cold = [row for row in rows if row["value"] < 30]
    max_value = max((row["value"] for row in rows), default=0)
    min_value = min((row["value"] for row in rows), default=0)
    dispersion = round(max_value - min_value, 1)
    state = "行业拥挤分化" if dispersion >= 55 else "行业温和分化" if dispersion >= 35 else "行业分布均衡"
    return {
        "available": True,
        "source": "market-breadth aggregated_data.json",
        "state": state,
        "detail": f"偏热行业 {len(hot)} 个，偏冷行业 {len(cold)} 个，行业宽度极差 {dispersion:.1f} 个百分点。",
        "hot_count": len(hot),
        "cold_count": len(cold),
        "dispersion": dispersion,
        "rows": sorted(rows, key=lambda row: row["value"], reverse=True),
    }


def _sentiment_extreme_payload(history_payload: dict) -> dict:
    history = history_payload.get("history") or []
    if not history:
        return {"state": "记录不足", "detail": "暂无历史情绪记录。"}
    latest = history[-1]
    scores = [row["score"] for row in history[-20:]]
    if len(scores) < 5:
        return {
            "state": "记录不足",
            "percentile": None,
            "detail": f"当前分数 {latest['score']}，本地仅有 {len(scores)} 条记录，暂不判断历史分位。",
        }
    percentile = round(sum(1 for score in scores if score <= latest["score"]) / len(scores) * 100)
    if latest["score"] >= 75:
        state = "过热"
    elif latest["score"] <= 30:
        state = "冰点"
    elif percentile >= 80:
        state = "近端高位"
    elif percentile <= 20:
        state = "近端低位"
    else:
        state = "常态区间"
    return {
        "state": state,
        "percentile": percentile,
        "detail": f"当前分数 {latest['score']}，处于近 {len(scores)} 条记录的约 {percentile}% 分位。",
    }


def _eastmoney_get(url: str, params: dict) -> dict | None:
    merged = {
        "ut": "7eea3edcaed734bea9cbfc24409ed989",
        **params,
    }
    try:
        response = requests.get(
            url,
            params=merged,
            headers={"User-Agent": "Mozilla/5.0", "Referer": "https://quote.eastmoney.com/"},
            timeout=15,
        )
        if response.status_code != 200:
            return None
        return response.json()
    except Exception:
        return None


def _record_sentiment(payload: dict) -> None:
    trade_date = str(payload.get("trade_time") or payload["generated_at"])[:10]
    rows = _load_sentiment_history()
    row = {
        "date": trade_date,
        "score": payload["score"],
        "label": payload["label"],
        "tag": payload["tag"],
        "generated_at": payload["generated_at"],
        "components": {item["key"]: item["score"] for item in payload["components"]},
    }
    rows = [item for item in rows if item.get("date") != trade_date]
    rows.append(row)
    rows = sorted(rows, key=lambda item: item["date"])[-120:]
    with SENTIMENT_HISTORY_PATH.open("w", encoding="utf-8") as file:
        json.dump(rows, file, ensure_ascii=False, indent=2)


def _load_sentiment_history() -> list[dict]:
    if not SENTIMENT_HISTORY_PATH.exists():
        return []
    with SENTIMENT_HISTORY_PATH.open(encoding="utf-8") as file:
        return json.load(file)


def _sentiment_history_payload(payload: dict) -> dict:
    history = _load_sentiment_history()
    latest = history[-1] if history else None
    previous = history[-2] if len(history) >= 2 else None
    recent = history[-20:]
    scores = [row["score"] for row in recent]
    alerts = []
    if latest:
        if latest["score"] >= 75:
            alerts.append("情绪进入过热区，需警惕拥挤交易和高位兑现。")
        elif latest["score"] <= 30:
            alerts.append("情绪进入偏冷区，若宽度不再恶化，后续可能出现超跌修复窗口。")
    if latest and previous:
        delta = latest["score"] - previous["score"]
        if abs(delta) >= 10:
            direction = "升温" if delta > 0 else "降温"
            alerts.append(f"情绪较上一记录明显{direction} {delta:+.0f} 分，短线交易状态发生变化。")
    if latest and scores and latest["score"] == max(scores) and len(scores) >= 5:
        alerts.append("当前情绪分数处于近 20 条记录高位。")
    if latest and scores and latest["score"] == min(scores) and len(scores) >= 5:
        alerts.append("当前情绪分数处于近 20 条记录低位。")
    return {
        "history": history,
        "latest": latest,
        "previous": previous,
        "recent_high": max(scores) if scores else payload["score"],
        "recent_low": min(scores) if scores else payload["score"],
        "alerts": alerts or ["暂无极值提醒，情绪处于可观察区间。"],
    }


def _average(values: list[float | None]) -> float | None:
    clean = [float(value) for value in values if value is not None]
    if not clean:
        return None
    return sum(clean) / len(clean)


def _clamp(value: float, low: float = 0, high: float = 100) -> float:
    return max(low, min(high, value))


def _latest_daily_brief_docx() -> Path | None:
    if not DAILY_BRIEF_OUTPUT_DIR.exists():
        return None
    files = sorted(DAILY_BRIEF_OUTPUT_DIR.glob("*.docx"), key=lambda path: path.stat().st_mtime, reverse=True)
    return files[0] if files else None


def _daily_brief_payload(docx_path: Path) -> dict:
    parsed = _parse_docx_brief(docx_path)
    return {
        "exists": True,
        "title": parsed["title"],
        "subtitle": parsed["subtitle"],
        "sections": parsed["sections"],
        "docx_path": str(docx_path),
        "download_url": "/api/daily-brief/download",
        "generated_at": datetime.fromtimestamp(docx_path.stat().st_mtime, TZ).isoformat(timespec="seconds"),
        "skill": str(DAILY_BRIEF_SKILL_DIR),
    }


def _fallback_daily_brief_payload(generated: bool = False) -> dict:
    generated_at = datetime.now(TZ).isoformat(timespec="seconds")

    def paragraph(text: str) -> dict:
        return {"type": "paragraph", "text": text, "style": "Normal"}

    def table(rows: list[list[str]]) -> dict:
        return {"type": "table", "rows": rows}

    sections = []
    conclusions = []
    market_rows = [["资产/模块", "状态", "核心观察"]]

    try:
        data_center = _safe_data_center_payload()
        coverage = data_center.get("coverage") or {}
        conclusions.append(
            f"云端数据库覆盖 {coverage.get('market_count', 0)} 类市场、{coverage.get('code_count', 0)} 个标的，最新日期 {coverage.get('max_date') or '-'}。"
        )
    except Exception as error:
        conclusions.append(f"数据中心暂不可用：{error}")

    try:
        cross = _safe_cross_market_risk_payload()
        summary = cross.get("summary") or {}
        conclusions.append(summary.get("conclusion") or "跨市场风险暂无结论。")
        market_rows.append(["跨市场风险", (cross.get("risk") or {}).get("state") or "-", "；".join((summary.get("analysis") or [])[:2])])
    except Exception as error:
        conclusions.append(f"跨市场风险暂不可用：{error}")

    try:
        macro = _macro_commodities_payload()
        summary = macro.get("summary") or {}
        conclusions.append(summary.get("conclusion") or "宏观商品暂无结论。")
        market_rows.append(["宏观商品", f"可用指标 {len([row for row in macro.get('items', []) if row.get('available', True)])}", "；".join((summary.get("analysis") or [])[:2])])
    except Exception as error:
        conclusions.append(f"宏观商品暂不可用：{error}")

    try:
        funds = _fund_mainline_payload("industry")
        summary = funds.get("summary") or {}
        conclusions.append(summary.get("conclusion") or "行业资金主线暂无结论。")
        leader = (funds.get("items") or [{}])[0]
        market_rows.append(["资金主线", leader.get("name") or "暂无主线", summary.get("conclusion") or "-"])
    except Exception as error:
        conclusions.append(f"资金主线暂不可用：{error}")

    breadth_stats = _breadth_stats(OUTPUT_DIR / "aggregated_data.json")
    if breadth_stats:
        conclusions.append(f"市场宽度均值 {breadth_stats['average']:.1f}%，最强行业为 {breadth_stats['strongest']['category']}，最弱行业为 {breadth_stats['weakest']['category']}。")
        market_rows.append(["市场宽度", f"{breadth_stats['average']:.1f}%", f"强：{breadth_stats['strongest']['category']}；弱：{breadth_stats['weakest']['category']}"])

    sections.append(
        {
            "heading": "摘要",
            "blocks": [
                paragraph("结论：" + (conclusions[0] if conclusions else "当前市场数据仍在准备中。")),
                *[paragraph(item) for item in conclusions[1:5]],
            ],
        }
    )
    sections.append(
        {
            "heading": "核心驱动",
            "blocks": [
                paragraph("当前简报按 daily-market-brief skill 的“先结论、后解释”原则生成，重点看 risk 偏好、资金主线、宏观压力和市场宽度是否共振。"),
                paragraph("若跨市场风险高、资金主线弱、市场宽度收缩，应优先降低高 beta 暴露；若资金主线连续且宽度扩散，则关注主线扩散和承接。"),
            ],
        }
    )
    sections.append(
        {
            "heading": "全球市场表现",
            "blocks": [
                table(market_rows),
                paragraph("股票指数、利率汇率、大宗商品和避险资产共同决定当日风险偏好；本页以现有云端接口可得数据为准，不补造缺失字段。"),
            ],
        }
    )
    sections.append(
        {
            "heading": "全球资金流向",
            "blocks": [
                paragraph("资金流向重点观察行业/概念主力净流入和阶段主线榜。若地域资金接口不可用，页面会降级显示行业资金流，避免用空白代替判断。"),
            ],
        }
    )
    sections.append(
        {
            "heading": "分类市场数据",
            "blocks": [
                paragraph("A 股、港股、美股和 VIX 数据优先来自云端 MySQL；宏观商品优先来自 FRED、akshare 和公开行情接口。"),
            ],
        }
    )
    sections.append(
        {
            "heading": "市场要闻",
            "blocks": [
                paragraph("当前云端版本只生成网页简报，不保存 Word 文件，也不下载文件；后续可接入新闻源后补充市场要闻。"),
            ],
        }
    )
    return {
        "ok": True,
        "exists": True,
        "title": "每日行情简报",
        "subtitle": f"生成时间 {generated_at[:10]}",
        "sections": sections,
        "docx_path": None,
        "download_url": None,
        "generated_at": generated_at,
        "skill": "daily-market-brief cloud renderer",
        "message": "已按 daily-market-brief 规则生成网页简报。" if generated else "当前显示云端网页简报。",
    }


def _parse_docx_brief(docx_path: Path) -> dict:
    from docx import Document

    document = Document(docx_path)
    title = docx_path.stem
    subtitle = ""
    sections: list[dict] = []
    current: dict | None = None

    for block in _iter_docx_blocks(document):
        if block["type"] == "paragraph":
            text = block["text"].strip()
            if not text:
                continue
            style = block.get("style", "")
            if style == "Title" or (not sections and current is None and title == docx_path.stem):
                title = text
                continue
            if not subtitle and text.startswith("交易日期"):
                subtitle = text
                continue
            if style.startswith("Heading 1"):
                current = {"heading": text, "blocks": []}
                sections.append(current)
                continue
            if current is None:
                current = {"heading": "导读", "blocks": []}
                sections.append(current)
            current["blocks"].append({"type": "paragraph", "text": text, "style": style})
        elif block["type"] == "table":
            if current is None:
                current = {"heading": "数据表", "blocks": []}
                sections.append(current)
            current["blocks"].append(block)

    return {"title": title, "subtitle": subtitle, "sections": sections}


def _iter_docx_blocks(document) -> list[dict]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    body = document.element.body
    blocks = []
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            blocks.append({"type": "paragraph", "text": paragraph.text, "style": paragraph.style.name})
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            blocks.append({"type": "table", "rows": rows})
    return blocks


def _breadth_stats(aggregated_path: Path) -> dict | None:
    if not aggregated_path.exists():
        return None
    with aggregated_path.open(encoding="utf-8") as file:
        payload = json.load(file)
    categories = payload.get("categories") or []
    dates = payload.get("dates") or []
    data = payload.get("data") or {}
    if not categories or not dates:
        return None

    latest_index = len(dates) - 1
    latest_values = []
    for category in categories:
        values = data.get(category) or []
        if latest_index < len(values) and values[latest_index] is not None:
            latest_values.append({"category": category, "value": float(values[latest_index])})
    if not latest_values:
        return None

    strongest = max(latest_values, key=lambda item: item["value"])
    weakest = min(latest_values, key=lambda item: item["value"])
    average = round(sum(item["value"] for item in latest_values) / len(latest_values), 1)
    improved = _change_rank(categories, data, latest_index, reverse=True)
    weakened = _change_rank(categories, data, latest_index, reverse=False)
    return {
        "latest_date": dates[-1],
        "dates_count": len(dates),
        "categories_count": len(categories),
        "average": average,
        "strongest": strongest,
        "weakest": weakest,
        "latest_values": latest_values,
        "improved": improved[:5],
        "weakened": weakened[:5],
    }


def _change_rank(categories: list[str], data: dict, latest_index: int, reverse: bool) -> list[dict]:
    rows = []
    lookback_index = max(0, latest_index - 5)
    for category in categories:
        values = data.get(category) or []
        if latest_index >= len(values) or lookback_index >= len(values):
            continue
        latest = values[latest_index]
        previous = values[lookback_index]
        if latest is None or previous is None:
            continue
        rows.append({"category": category, "change": round(float(latest) - float(previous), 1), "value": float(latest)})
    return sorted(rows, key=lambda item: item["change"], reverse=reverse)


def _secondary_industry_data(raw_path: Path) -> dict:
    if not raw_path.exists() or not INDUSTRY_CATEGORY_MAP:
        return {}
    with raw_path.open(encoding="utf-8") as file:
        raw = json.load(file)
    industries = raw.get("industries") or []
    dates = raw.get("dates") or []
    data_points = raw.get("data") or []
    children: dict[str, dict[str, list[float | None]]] = {}

    for industry in industries:
        category = INDUSTRY_CATEGORY_MAP.get(industry)
        if category:
            children.setdefault(category, {})[industry] = [None for _ in dates]

    for date_idx, industry_idx, value in data_points:
        if industry_idx >= len(industries) or date_idx >= len(dates):
            continue
        industry = industries[industry_idx]
        category = INDUSTRY_CATEGORY_MAP.get(industry)
        if not category:
            continue
        children.setdefault(category, {}).setdefault(industry, [None for _ in dates])[date_idx] = (
            round(float(value), 1) if value and value > 0 else None
        )

    return {
        category: dict(sorted(industry_rows.items(), key=lambda item: _latest_value(item[1]), reverse=True))
        for category, industry_rows in children.items()
    }


def _latest_value(values: list[float | None]) -> float:
    for value in reversed(values):
        if value is not None:
            return float(value)
    return -1.0


def _market_commentary(stats: dict | None) -> dict | None:
    if not stats:
        return None

    average = stats["average"]
    strongest = stats["strongest"]
    weakest = stats["weakest"]
    latest_values = stats.get("latest_values") or []
    improved = stats.get("improved") or []
    weakened = stats.get("weakened") or []
    strong_count = sum(1 for item in latest_values if item["value"] >= 60)
    weak_count = sum(1 for item in latest_values if item["value"] < 30)

    if average >= 60:
        conclusion = "市场宽度处于偏强区间，行业情绪整体活跃。"
    elif average >= 45:
        conclusion = "市场宽度处于中性偏稳区间，行业情绪分化但仍有结构性机会。"
    elif average >= 30:
        conclusion = "市场宽度偏弱，行业情绪以防守和收缩为主。"
    else:
        conclusion = "市场宽度处于低位，行业情绪明显偏弱。"

    analysis = [
        f"截至 {stats['latest_date']}，26 个一级行业 MA20 站上率均值为 {average:.1f}%。",
        f"强势方向集中在 {strongest['category']}，最新站上率 {strongest['value']:.1f}%；弱势方向为 {weakest['category']}，仅 {weakest['value']:.1f}%。",
        f"站上率超过 60% 的行业有 {strong_count} 个，低于 30% 的行业有 {weak_count} 个，市场内部广度仍需观察。",
    ]
    if improved:
        leader = improved[0]
        analysis.append(f"近 5 日相对改善最明显的是 {leader['category']}，变化 {leader['change']:+.1f} 个百分点。")
    if weakened:
        laggard = weakened[0]
        analysis.append(f"近 5 日走弱最明显的是 {laggard['category']}，变化 {laggard['change']:+.1f} 个百分点。")

    return {"conclusion": conclusion, "analysis": analysis}
